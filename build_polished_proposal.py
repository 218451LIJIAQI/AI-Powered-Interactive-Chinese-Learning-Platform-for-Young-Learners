"""Build the polished FYP proposal as a properly formatted Word document.

This script produces an APA-style FYP proposal aligned with the official
template prepared by Dr. Azrina Kamaruddin (SEIS Project Coordinator
Department), Universiti Putra Malaysia. The script avoids placing
implementation-specific claims in past tense and writes from the
perspective of a student who is proposing the project.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


# ---------------------------------------------------------------------------
# Document configuration helpers
# ---------------------------------------------------------------------------

OUTPUT_FILENAME = "Li_Jiaqi_FYP_Proposal_Polished_Final.docx"
PROJECT_ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = PROJECT_ROOT / OUTPUT_FILENAME


def _set_cell_borders(cell):
    """Apply thin black borders to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "000000")
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def _set_cell_shading(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _add_page_number_field(paragraph) -> None:
    """Insert a PAGE field into the given paragraph."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE \\* MERGEFORMAT"
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


BODY_FONT = "Arial"


def _set_run_font(run, name: str = BODY_FONT) -> None:
    """Force the East Asian and ASCII font hints so Word respects the choice."""
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _configure_styles(doc: Document) -> None:
    """Configure default styles per the SEIS Proposal Report Format."""
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)

    for level in range(1, 4):
        style = styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        style.paragraph_format.space_before = Pt(18)
        style.paragraph_format.space_after = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.keep_with_next = True
        if level == 1:
            # Chapter label uses larger title-case Arial heading per spec; the
            # following 4-line space rule is enforced via space_after.
            style.font.size = Pt(18)
            style.paragraph_format.space_after = Pt(36)
        elif level == 2:
            style.font.size = Pt(14)
            style.paragraph_format.space_after = Pt(12)
        else:
            style.font.size = Pt(12)
            style.paragraph_format.space_after = Pt(12)
        h_rpr = style.element.get_or_add_rPr()
        h_rfonts = h_rpr.find(qn("w:rFonts"))
        if h_rfonts is None:
            h_rfonts = OxmlElement("w:rFonts")
            h_rpr.append(h_rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            h_rfonts.set(qn(attr), BODY_FONT)


def _configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_page_number_field(p)
    for run in p.runs:
        _set_run_font(run, BODY_FONT)
        run.font.size = Pt(12)


# ---------------------------------------------------------------------------
# Convenience writers
# ---------------------------------------------------------------------------

def add_paragraph(doc: Document, text: str, *, italic: bool = False,
                  alignment=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.first_line_indent = Cm(1.27)
    run = p.add_run(text)
    run.italic = italic
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(12)


def add_plain_paragraph(doc: Document, text: str, *,
                        alignment=WD_ALIGN_PARAGRAPH.LEFT,
                        bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(12)


def add_heading(doc: Document, text: str, level: int) -> None:
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        _set_run_font(run, BODY_FONT)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(12)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(12)


def add_table(doc: Document, header: list[str], rows: list[list[str]],
              *, header_fill: str = "D9E1F2",
              column_widths_cm: list[float] | None = None) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"

    for idx, label in enumerate(header):
        cell = table.cell(0, idx)
        cell.text = ""
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(label)
        run.bold = True
        _set_run_font(run, BODY_FONT)
        run.font.size = Pt(11)
        _set_cell_shading(cell, header_fill)
        _set_cell_borders(cell)

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            para = cell.paragraphs[0]
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(value)
            _set_run_font(run, BODY_FONT)
            run.font.size = Pt(11)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_borders(cell)

    if column_widths_cm:
        for row in table.rows:
            for c_idx, width in enumerate(column_widths_cm):
                row.cells[c_idx].width = Cm(width)

    # SEIS spec: 2 spacings after a table before next content
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_figure_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(11)


def add_table_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(11)


def add_reference(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(12)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


# ---------------------------------------------------------------------------
# Section 0 - Title page and front matter
# ---------------------------------------------------------------------------

def _add_centered(doc: Document, text: str, *, size: int = 14,
                  bold: bool = False, italic: bool = False,
                  space_before: int = 0, space_after: int = 0) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    _set_run_font(run, BODY_FONT)
    run.font.size = Pt(size)


def write_title_page(doc: Document) -> None:
    # Title block (Arial, Title Case, max 12 words, no contractions/abbreviations)
    _add_centered(
        doc,
        "Artificial Intelligence-Powered Interactive Chinese "
        "Learning Platform for Young Learners in Malaysia",
        size=22,
        bold=True,
        space_before=120,
        space_after=12,
    )
    _add_centered(doc, "Proposal Report", size=16, bold=True, space_after=24)

    # Spacer between title block and student information
    spacer1 = doc.add_paragraph()
    spacer1.paragraph_format.space_before = Pt(96)

    _add_centered(doc, "Li Jiaqi", size=14, bold=True)
    _add_centered(doc, "218451", size=14, bold=True, space_after=24)

    # Spacer between student information and programme block
    spacer2 = doc.add_paragraph()
    spacer2.paragraph_format.space_before = Pt(72)

    _add_centered(doc, "Bachelor of Software Engineering with Honours",
                  size=12)
    _add_centered(
        doc,
        "Department of Software Engineering and Information System",
        size=12,
    )
    _add_centered(
        doc,
        "Faculty of Computer Science and Information Technology",
        size=12,
    )
    _add_centered(doc, "Universiti Putra Malaysia", size=12, space_after=24)

    spacer3 = doc.add_paragraph()
    spacer3.paragraph_format.space_before = Pt(24)

    _add_centered(doc, "Semester 2, 2025/2026", size=12)
    _add_centered(doc, "Supervisor: Dr. Azrina binti Kamaruddin",
                  size=12, italic=True)

    add_page_break(doc)


def write_declaration_page(doc: Document) -> None:
    add_heading(doc, "Declaration of Originality", 1)
    add_paragraph(
        doc,
        "I, Li Jiaqi (Matric No. 218451), hereby declare that this Final Year "
        "Project proposal entitled \"Artificial Intelligence-Powered "
        "Interactive Chinese Learning Platform for Young Learners in "
        "Malaysia\" is the result of my own research and original effort, "
        "except for citations and quotations that have been duly "
        "acknowledged. I further declare that this proposal has not been "
        "previously or concurrently submitted in candidature for any other "
        "degree at Universiti Putra Malaysia or any other institution. I "
        "understand that any form of academic misconduct may lead to "
        "disciplinary action under the academic regulations of Universiti "
        "Putra Malaysia.",
    )

    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(48)
    sig.add_run("Signature: ____________________________").bold = False

    name = doc.add_paragraph()
    name.add_run("Name: Li Jiaqi").bold = False

    matric = doc.add_paragraph()
    matric.add_run("Matric Number: 218451").bold = False

    date = doc.add_paragraph()
    date.add_run("Date: 13 May 2026").bold = False

    add_page_break(doc)


def write_acknowledgement(doc: Document) -> None:
    add_heading(doc, "Acknowledgement", 1)
    add_paragraph(
        doc,
        "I would like to express my deepest gratitude to my supervisor, "
        "Dr. Azrina binti Kamaruddin, for her continuous guidance, constructive "
        "feedback, and unwavering support throughout the formulation of this "
        "Final Year Project proposal. Her insights have been instrumental in "
        "shaping both the academic rigour and the pedagogical orientation of "
        "this work.",
    )
    add_paragraph(
        doc,
        "I am sincerely thankful to the Department of Software Engineering and "
        "Information System and the Faculty of Computer Science and "
        "Information Technology, Universiti Putra Malaysia, for providing the "
        "academic environment, library resources, and laboratory facilities "
        "that have enabled me to undertake this project.",
    )
    add_paragraph(
        doc,
        "I would also like to thank the Mandarin teachers and educators "
        "in Malaysia who have offered general guidance and encouragement "
        "while I was conceptualising this proposal; their perspectives "
        "have helped me to keep the proposed design grounded in real "
        "classroom and home-learning realities. Finally, I extend my "
        "appreciation to my family, course mates, and friends for their "
        "continuous encouragement and patience throughout the proposal "
        "preparation period.",
    )
    add_page_break(doc)


def write_abstract(doc: Document) -> None:
    add_heading(doc, "Abstract", 1)
    add_paragraph(
        doc,
        "Mandarin Chinese is increasingly important to Malaysian young learners "
        "due to the country's multilingual classroom culture, the prominence of "
        "Sekolah Jenis Kebangsaan Cina (SJKC) and Mandarin enrichment "
        "programmes, and the economic relevance of the language across "
        "ASEAN. However, beginner Mandarin learners aged 7 to 12 face "
        "compounded difficulties: they must simultaneously acquire Hanyu "
        "Pinyin, four lexical tones, vocabulary, listening comprehension, "
        "introductory sentence patterns, and Chinese characters, while "
        "navigating fragmented digital tools, limited formative feedback, and "
        "heavy teacher workload. Existing platforms such as Duolingo, "
        "HelloChinese, ChineseSkill, LingoDeer, and Skritter remain primarily "
        "self-paced, are not designed for the Malaysian multilingual context, "
        "and rarely integrate teacher-managed workflows with bounded AI "
        "support that is age-appropriate.",
    )
    add_paragraph(
        doc,
        "This proposal presents the planned design, development, and "
        "evaluation of a web-based Artificial Intelligence (AI)-Powered "
        "Interactive Chinese Learning Platform for Young Learners. The "
        "proposed prototype integrates structured beginner Mandarin "
        "content, gamified interactive activities (flashcards, quizzes, "
        "matching games, listening tasks, and introductory tracing), "
        "bounded AI assistance through an external OpenAI-compatible "
        "API, role-based access control with teacher verification, "
        "multilingual interfaces in English, Bahasa Melayu, and "
        "Simplified Chinese, recorded and live learning support, and "
        "basic moderation, support, and analytics workflows for Student, "
        "Teacher, Administrator, and Guest users.",
    )
    add_paragraph(
        doc,
        "The study adopts Design and Development Research (DDR) as its "
        "overall research design, complemented by Agile incremental "
        "development across six increments. Data collection includes "
        "literature review, expert consultation with Mandarin educators, "
        "task-based usability testing, a Likert-scale questionnaire informed by "
        "the System Usability Scale (SUS), and observational notes during "
        "supervised testing sessions. The proposed full-stack prototype is "
        "planned to use Vue 3, TypeScript, Element Plus, and Vue I18n on the "
        "client side; Spring Boot 3.5.x, Java 21, Spring Security, JWT, and "
        "Spring Data JPA on the server side; and MySQL 8.4 LTS as the "
        "relational data store, orchestrated locally via Docker Compose.",
    )
    add_paragraph(
        doc,
        "Evaluation will examine technical feasibility, usability (SUS targets "
        "of 68 or higher), perceived learning support, and child-appropriate "
        "design through expert review and supervised user testing. The "
        "expected contributions are fourfold: (i) an academically grounded "
        "design model that connects beginner Mandarin learning challenges "
        "with Cognitive Load Theory, the Zone of Proximal Development, "
        "motivation theory, and bounded AI ethics; (ii) a reusable full-stack "
        "prototype architecture for Malaysian educational technology projects; "
        "(iii) evaluation instruments for child-oriented language-learning "
        "platforms; and (iv) a feasibility demonstration that AI in education "
        "can be deployed responsibly and ethically when paired with teacher "
        "oversight and culturally responsive design.",
    )
    p = doc.add_paragraph()
    run = p.add_run("Keywords: ")
    run.bold = True
    p.add_run(
        "Mandarin learning, young learners, educational technology, "
        "AI-assisted learning, gamification, scaffolding, multilingual "
        "platform, Malaysia, full-stack web application, Final Year Project."
    )
    add_page_break(doc)


def write_table_of_contents(doc: Document) -> None:
    add_heading(doc, "Table of Contents", 1)
    add_paragraph(
        doc,
        "The Table of Contents will be auto-generated in the final submission "
        "using Microsoft Word's References tab. To update it, position the "
        "cursor on the next page in Word, open References, click Table of "
        "Contents, then choose an automatic style. The placeholder structure "
        "below mirrors the template prescribed by the SEIS Project "
        "Coordinator Department.",
        italic=True,
    )

    headings = [
        "Declaration of Originality",
        "Acknowledgement",
        "Abstract",
        "List of Tables",
        "List of Figures",
        "List of Abbreviations",
        "1.0 Introduction",
        "    1.1 Background of the Study",
        "    1.2 Motivation of the Study",
        "    1.3 Project Overview",
        "2.0 Problem Statement",
        "    2.1 Problem Identification",
        "    2.2 Problem Justification",
        "    2.3 Impact of the Problem",
        "3.0 Objectives",
        "    3.1 General Objective",
        "    3.2 Specific Objectives",
        "4.0 Scope of the Study",
        "    4.1 System Scope",
        "    4.2 User Scope",
        "    4.3 Dataset Scope",
        "5.0 Literature Review",
        "    5.1 Introduction to Literature Review",
        "    5.2 Related Works / Previous Studies",
        "    5.3 Theoretical Framework / Models",
        "    5.4 Research Gap Analysis",
        "    5.5 Summary",
        "6.0 Methodology",
        "    6.1 Introduction",
        "    6.2 Research Design",
        "    6.3 Development Methodology",
        "    6.4 Data Collection Methods",
        "    6.5 System Design (Architecture, Tools, Technologies)",
        "    6.6 Software Verification and Validation",
        "    6.7 Evaluation Methods",
        "    6.8 Summary",
        "7.0 Return on Investment / Return on Value",
        "    7.1 Cost-Benefit Analysis",
        "    7.2 Expected Benefits (Tangible and Intangible)",
        "    7.3 Value Contribution to Stakeholders",
        "8.0 Project Planning",
        "    8.1 Work Breakdown Structure",
        "    8.2 Gantt Chart",
        "9.0 Conclusion",
        "    9.1 Summary of Proposal",
        "    9.2 Expected Contributions",
        "10.0 References",
        "11.0 Appendices",
        "    Appendix A. Project Timeline Details",
        "    Appendix B. Survey / Interview Instruments",
        "    Appendix C. Additional Diagrams",
    ]
    for heading in headings:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = para.add_run(heading)
        _set_run_font(run, BODY_FONT)
        run.font.size = Pt(12)
    add_page_break(doc)


def write_lists_and_abbreviations(doc: Document) -> None:
    add_heading(doc, "List of Tables", 1)
    tables_list = [
        "Table 5.1. Critical comparison of selected existing Chinese learning systems.",
        "Table 5.2. Research gap analysis matrix between existing solution categories and the proposed platform.",
        "Table 6.1. Mapping of specific objectives to research design stages and evaluation evidence.",
        "Table 6.2. Proposed tools and technologies.",
        "Table 6.3. Evaluation criteria, instruments, and acceptance thresholds.",
        "Table 7.1. Estimated cost structure for the proposed prototype (in MYR).",
        "Table 7.2. Expected tangible and intangible benefits of the proposed project.",
        "Table 8.1. Proposed Work Breakdown Structure (WBS) mapped to the methodology.",
        "Table 8.2. Risk register and mitigation plan.",
    ]
    for entry in tables_list:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = para.add_run(entry)
        _set_run_font(run, BODY_FONT)
        run.font.size = Pt(12)

    add_heading(doc, "List of Figures", 1)
    figures_list = [
        "Figure 1.1. Conceptual overview of the proposed Artificial Intelligence-Powered Interactive Chinese Learning Platform.",
        "Figure 5.1. Proposed conceptual framework linking beginner Mandarin learning challenges, platform components, and expected outcomes.",
        "Figure 6.1. Proposed methodology combining Design and Development Research with Agile incremental development.",
        "Figure 6.2. Proposed three-tier system architecture with role-based access and external integrations.",
        "Figure 8.1. Proposed Final Year Project Gantt chart spanning 24 weeks.",
    ]
    for entry in figures_list:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = para.add_run(entry)
        _set_run_font(run, BODY_FONT)
        run.font.size = Pt(12)

    add_heading(doc, "List of Abbreviations", 1)
    abbreviations = [
        ("AI", "Artificial Intelligence"),
        ("API", "Application Programming Interface"),
        ("CFL", "Chinese as a Foreign Language"),
        ("CSL", "Chinese as a Second Language"),
        ("CLT", "Cognitive Load Theory"),
        ("DDR", "Design and Development Research"),
        ("FYP", "Final Year Project"),
        ("FSKTM", "Faculty of Computer Science and Information Technology"),
        ("JWT", "JSON Web Token"),
        ("KSSR", "Kurikulum Standard Sekolah Rendah"),
        ("LMS", "Learning Management System"),
        ("MOE", "Ministry of Education Malaysia"),
        ("MYR", "Malaysian Ringgit"),
        ("PDPA", "Personal Data Protection Act 2010 (Malaysia)"),
        ("RBAC", "Role-Based Access Control"),
        ("ROI", "Return on Investment"),
        ("ROV", "Return on Value"),
        ("SDLC", "Software Development Life Cycle"),
        ("SEIS", "Software Engineering and Information System"),
        ("SJKC", "Sekolah Jenis Kebangsaan Cina"),
        ("SPA", "Single-Page Application"),
        ("SRS", "Software Requirements Specification"),
        ("SUS", "System Usability Scale"),
        ("UAT", "User Acceptance Testing"),
        ("UI", "User Interface"),
        ("UPM", "Universiti Putra Malaysia"),
        ("UX", "User Experience"),
        ("WBS", "Work Breakdown Structure"),
        ("ZPD", "Zone of Proximal Development"),
    ]
    add_table(
        doc,
        header=["Abbreviation", "Meaning"],
        rows=[[abbr, meaning] for abbr, meaning in abbreviations],
        column_widths_cm=[4.0, 12.0],
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 1 - Introduction
# ---------------------------------------------------------------------------

def write_introduction(doc: Document) -> None:
    add_heading(doc, "1.0 Introduction", 1)
    add_heading(doc, "1.1 Background of the Study", 2)
    add_paragraph(
        doc,
        "Mandarin Chinese is one of the most widely spoken languages in "
        "the world and a strategically important language for Malaysia's "
        "economic, educational, and cultural engagement with East Asia. "
        "Within the Malaysian primary education system, Mandarin is a "
        "compulsory subject in Sekolah Jenis Kebangsaan Cina (SJKC) and "
        "is also offered as an elective in Sekolah Kebangsaan, private "
        "and international schools, tuition centres, and the Bahasa Cina "
        "Sekolah Kebangsaan component of the Kurikulum Standard Sekolah "
        "Rendah (KSSR Semakan 2017) (Kementerian Pendidikan Malaysia, "
        "2017). The co-existence of Bahasa Melayu, English, Mandarin, "
        "and various community languages creates a uniquely multilingual "
        "classroom environment in which pupils aged 7 to 12 may "
        "simultaneously learn multiple writing systems, phonologies, and "
        "grammatical patterns (Chua et al., 2024; Rahmat et al., 2021).",
    )
    add_paragraph(
        doc,
        "Despite this exposure, Mandarin remains particularly demanding for "
        "beginners because it requires the integrated acquisition of Hanyu "
        "Pinyin, four lexical tones, basic vocabulary, morpho-syllabic Chinese "
        "characters, listening comprehension, and introductory sentence "
        "patterns (Rahmat et al., 2021; Wang, 2024). Recent cognitive "
        "research highlights that working memory capacity is limited when "
        "learners encounter unfamiliar material, and that overloading novice "
        "learners can produce confusion, frustration, or learning failure "
        "(Sweller et al., 2019). Butler (2022) further notes that primary-aged "
        "learners are still developing executive function, theory of mind, and "
        "metalinguistic awareness, meaning instructional materials designed "
        "for adults cannot be deployed for children without careful adaptation. "
        "These cognitive considerations have direct implications for any "
        "Mandarin learning environment intended for the 7 to 12 age group in "
        "Malaysia.",
    )
    add_paragraph(
        doc,
        "Concurrently, digital technologies are rapidly transforming language "
        "education in Malaysia and the wider region. Liu (2018) describes the "
        "growth of online, blended, and residential Chinese instruction, while "
        "Yang (2020) reports increasing parental investment in early online "
        "language learning. Hongsamsibkao (2024) demonstrates that a "
        "structured online Chinese course can improve listening "
        "comprehension and Pinyin pronunciation when designed with explicit "
        "scaffolding. Within Malaysia, Rahmat et al. (2021) found that online "
        "interaction promotes engagement among Mandarin learners by "
        "facilitating richer interaction with teachers, peers, and content. Chua "
        "et al. (2024), comparing Malaysian and German contexts, report that "
        "technology integration is promising provided that teacher support and "
        "learner motivation are preserved. Fang et al. (2024) likewise "
        "synthesise mobile-assisted Chinese language learning research and "
        "conclude that integrated, beginner-oriented platforms remain "
        "comparatively scarce.",
    )
    add_paragraph(
        doc,
        "Artificial intelligence (AI) is the most recent technological "
        "development reshaping educational technology. Systematic reviews by "
        "Wang et al. (2024) and Garz\u00f3n et al. (2025) report rapid adoption "
        "of AI for intelligent tutoring, adaptive feedback, content generation, "
        "and learning analytics, while raising ethical concerns related to "
        "teacher readiness, data privacy, and overdependence. In language "
        "learning specifically, Du and Daniel (2024) found that AI-powered "
        "chatbots improved English speaking practice by lowering the "
        "affective barrier of practice and increasing engagement. Translating "
        "these findings into the Malaysian Mandarin context requires careful "
        "boundary-setting because young learners are particularly sensitive "
        "to safety, interface complexity, and the appropriateness of "
        "AI-generated content.",
    )
    add_paragraph(
        doc,
        "Against this background, this study proposes the design, development, "
        "and evaluation of a web-based AI-Powered Interactive Chinese "
        "Learning Platform tailored for beginner Mandarin learners aged 7 to "
        "12 within the Malaysian multilingual setting. The proposed platform is "
        "intended to combine child-centred pedagogy, scaffolded learning "
        "experiences, multimodal interactive activities, bounded AI assistance, "
        "and basic teacher-managed workflows in one coherent FYP-scale "
        "prototype. The remainder of Chapter 1 outlines the motivation for the "
        "study and provides an overview of the proposed project.",
    )

    add_heading(doc, "1.2 Motivation of the Study", 2)
    add_paragraph(
        doc,
        "The motivation for this study is grounded in four complementary "
        "considerations: pedagogical, motivational, technological, and "
        "personal-academic.",
    )
    add_paragraph(
        doc,
        "First, beginner Mandarin learning in Malaysia is pedagogically "
        "demanding because learners must concurrently master Pinyin, tones, "
        "vocabulary, listening, sentence structure, and Chinese characters. "
        "Paee et al. (2024), studying Malaysian tertiary Mandarin learners, "
        "report that demotivating factors include assessment pressure, "
        "content difficulty, perceived workload, boredom, limited support, and "
        "the difficulty of Chinese characters. Although their participants were "
        "older than the target population of this proposal, the literature "
        "suggests that younger beginners are likely to face equivalent or even "
        "stronger demotivators because they possess less developed "
        "self-regulation and metacognitive strategies (Butler, 2022; Sweller "
        "et al., 2019). A platform that decomposes content into shorter, "
        "scaffolded steps and provides regular encouragement is therefore "
        "pedagogically justified.",
    )
    add_paragraph(
        doc,
        "Second, motivational research consistently shows that young "
        "learners thrive when learning is enjoyable, structured, and supported "
        "by encouraging teachers. Wallace and Leong (2020) document that "
        "primary EFL learners are particularly motivated by games, songs, and "
        "supportive teacher relationships. Roorda et al. (2017), in a "
        "meta-analysis of 189 studies covering approximately 250,000 "
        "students, demonstrate that affective teacher\u2013student "
        "relationships are positively associated with engagement and "
        "achievement. These findings suggest that integrating gamified "
        "practice and lightweight teacher communication into a Mandarin "
        "learning prototype is likely to mitigate motivational decline among "
        "young beginners.",
    )
    add_paragraph(
        doc,
        "Third, AI technology now enables realistic, low-cost integration of "
        "intelligent learning supports without the overhead of training a "
        "custom model. External AI APIs accessible from Malaysia\u2014such "
        "as DeepSeek through an OpenAI-compatible interface\u2014can "
        "provide bounded forms of support including short explanations, "
        "example sentences, vocabulary hints, beginner-friendly feedback, "
        "and teacher material drafting. By treating AI as a supportive tool "
        "rather than an autonomous teacher or assessment authority, the "
        "proposed prototype aligns with the cautionary literature on AI in "
        "education (Garz\u00f3n et al., 2025; Wang et al., 2024) and remains "
        "ethically appropriate for child-oriented use.",
    )
    add_paragraph(
        doc,
        "Fourth, from a personal-academic perspective, the project enables a "
        "Bachelor of Software Engineering candidate at Universiti Putra "
        "Malaysia to integrate the technical, design, and research "
        "competencies cultivated across the programme. The prototype "
        "involves authentication, role-based access control, content "
        "management, third-party API integration, gamified user experience "
        "design, and basic analytics\u2014all of which are core software "
        "engineering competencies aligned with the SEIS Department's "
        "graduate attributes. The project therefore offers an opportunity to "
        "translate research on young learners, motivation, scaffolding, and "
        "digital learning into a practically grounded Mandarin learning "
        "prototype that is both academically rigorous and locally relevant.",
    )

    add_heading(doc, "1.3 Project Overview", 2)
    add_paragraph(
        doc,
        "The proposed project, hereafter referred to as the Artificial "
        "Intelligence (AI)-Powered Interactive Chinese Learning Platform "
        "for Young Learners, will be designed, developed, and evaluated as "
        "a full-stack web prototype for children aged 7 to 12 who are "
        "beginning to learn Mandarin in the Malaysian multilingual "
        "environment. The proposed system will support four user "
        "categories: Student, Teacher, Administrator, and Guest. A "
        "simplified conceptual overview of the proposed platform is shown "
        "in Figure 1.1.",
    )
    add_figure_caption(
        doc,
        "Figure 1.1. Conceptual overview of the proposed Artificial "
        "Intelligence-Powered Interactive Chinese Learning Platform, "
        "illustrating how structured beginner Mandarin content, gamified "
        "interactive activities, bounded AI support, teacher-managed "
        "workflows, and administrator oversight are brought together for "
        "Malaysian young learners.",
    )
    add_paragraph(
        doc,
        "For students, the proposed prototype will deliver structured "
        "beginner-level content covering Pinyin, basic vocabulary, sentence "
        "patterns, listening comprehension, and introductory Chinese "
        "character practice. Interactive activities will include flashcards, "
        "multiple-choice quizzes, matching games, listening exercises, and "
        "introductory tracing or writing tasks. To encourage regular practice, "
        "reward mechanisms such as points, badges, levels, daily streaks, "
        "and achievement tracking will be integrated, while AI-assisted "
        "features will be limited to supportive, non-evaluative functions such "
        "as short explanations, vocabulary hints, sentence practice prompts, "
        "and beginner-friendly feedback.",
    )
    add_paragraph(
        doc,
        "For teachers, the proposed prototype will offer tools for creating and "
        "managing learning content, drafting materials with AI assistance, "
        "monitoring student progress, and providing structured feedback. "
        "Teacher accounts will be subject to a verification workflow before "
        "they are permitted to publish or moderate learning content, ensuring "
        "that public-facing material is curated by qualified educators. For "
        "administrators, the proposed system will provide user management, "
        "teacher verification, content moderation, support and report handling, "
        "audit logging, and basic analytics dashboards.",
    )
    add_paragraph(
        doc,
        "Recorded learning will be supported through uploaded or linked "
        "video and audio resources, while live learning coordination will be "
        "handled through external meeting links such as Google Meet or "
        "Zoom rather than an in-house video conferencing platform. Payment "
        "and subscription workflows will be implemented as sandbox-friendly "
        "simulations to demonstrate the commerce flow without processing "
        "real financial transactions, in line with FYP scope and Malaysian "
        "consumer protection considerations.",
    )
    add_paragraph(
        doc,
        "Technically, the proposed prototype will be built using Vue 3, Vite, "
        "TypeScript, Pinia, Vue Router, Axios, Vue I18n, and Element Plus on "
        "the frontend; Spring Boot 3.5.x with Java 21, Spring Security, JWT, "
        "Spring Data JPA, and Maven on the backend; and MySQL 8.4 LTS as "
        "the relational data store. Local infrastructure will be orchestrated "
        "via Docker Compose, and AI-assisted features will be integrated "
        "through a configurable external API behind a backend proxy that "
        "protects credentials. Multilingual interfaces will be provided in "
        "English, Bahasa Melayu, and Simplified Chinese to align with the "
        "Malaysian linguistic landscape.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 2 - Problem Statement
# ---------------------------------------------------------------------------

def write_problem_statement(doc: Document) -> None:
    add_heading(doc, "2.0 Problem Statement", 1)
    add_heading(doc, "2.1 Problem Identification", 2)
    add_paragraph(
        doc,
        "Although digital Chinese learning resources are increasingly "
        "available in Malaysia, several issues continue to limit their "
        "effectiveness for beginner learners aged 7 to 12 in a multilingual "
        "context. Five interrelated issues are identified below, each grounded "
        "in current literature and informal observation of local practice.",
    )
    add_paragraph(
        doc,
        "Issue 1: Cognitive overload from simultaneous skill demands. Beginner "
        "Mandarin requires the parallel acquisition of Hanyu Pinyin, four "
        "lexical tones, basic vocabulary, character recognition, character "
        "writing, and sentence structure. This combination is challenging for "
        "children whose working memory capacity is still developing because "
        "the volume of unfamiliar information may exceed cognitive limits "
        "(Sweller et al., 2019). Paee et al. (2024) likewise report that content "
        "difficulty and the demands of character learning are leading "
        "demotivators among Malaysian Mandarin learners. Many existing "
        "consumer apps present learners with rapid skill switches and "
        "minimal scaffolding, intensifying overload risk for the 7 to 12 age "
        "group.",
    )
    add_paragraph(
        doc,
        "Issue 2: Limited formative feedback in pronunciation, tone, and "
        "writing practice. Although structured online Chinese instruction has "
        "been shown to improve listening and Pinyin pronunciation when "
        "appropriately designed (Chen, 2024; Hongsamsibkao, 2024), many "
        "Malaysian learners rely on static notes, prerecorded videos, or "
        "drill-based exercises that offer little corrective guidance when a "
        "learner mispronounces a tone or mis-strokes a character. The "
        "absence of timely, age-appropriate feedback can allow errors to "
        "fossilise during the critical beginner phase.",
    )
    add_paragraph(
        doc,
        "Issue 3: Fragmentation of digital learning tools. Local learners "
        "frequently use one application for vocabulary practice, another for "
        "Pinyin learning, another for handwriting, and separate channels for "
        "teacher communication. Fang et al. (2024) note that mobile-assisted "
        "Chinese learning research often focuses on isolated skills rather than "
        "integrated support. This fragmentation interrupts learning continuity, "
        "increases cognitive switching costs for children, and complicates "
        "progress visibility for teachers, parents, and tuition centres.",
    )
    add_paragraph(
        doc,
        "Issue 4: Excessive teacher workload for beginner support. Mandarin "
        "teachers in Malaysian SJKC, Sekolah Kebangsaan POL classes, "
        "and tuition centres often invest significant effort in preparing "
        "beginner-friendly materials, marking practice items, and providing "
        "individual feedback. Roorda et al. (2017) emphasise that teacher "
        "support is critical for engagement and achievement; however, when "
        "much of the teacher's time is consumed by routine content "
        "preparation and progress checking, less time remains for the higher "
        "value coaching that young learners need most.",
    )
    add_paragraph(
        doc,
        "Issue 5: Risks associated with general-purpose AI for young learners. "
        "Mainstream AI assistants are typically designed for adult productivity "
        "and may produce content that is linguistically advanced, culturally "
        "inappropriate, or insufficiently moderated for primary-aged users. "
        "Garz\u00f3n et al. (2025) and Wang et al. (2024) caution that AI in "
        "education raises ethical, accuracy, teacher readiness, and "
        "overdependence concerns, all of which are amplified when the user "
        "is a child. There is therefore a need for a bounded, supervised, and "
        "child-appropriate model of AI integration in language learning.",
    )
    add_paragraph(
        doc,
        "Considered together, these five issues reveal a clear practical need: "
        "an integrated, child-appropriate, feedback-rich, and teacher-managed "
        "web platform that combines beginner Mandarin content, AI-assisted "
        "support, gamified practice, multilingual interfaces, and moderation "
        "controls in one coherent prototype tailored to the Malaysian context. "
        "Within the scope of the literature reviewed in Chapter 5, no existing "
        "system or research artefact appears to address all of these "
        "dimensions simultaneously for the 7 to 12 age group.",
    )

    add_heading(doc, "2.2 Problem Justification", 2)
    add_paragraph(
        doc,
        "The need for the proposed platform can be justified on cognitive, "
        "motivational, linguistic, technological, and contextual grounds.",
    )
    add_paragraph(
        doc,
        "Cognitively, Cognitive Load Theory (Sweller et al., 2019) provides a "
        "principled rationale for sequencing beginner Mandarin content into "
        "shorter, structured tasks with limited extraneous load. Switching "
        "between unrelated apps for Pinyin, vocabulary, tones, and characters "
        "is likely to add extraneous cognitive load without contributing to "
        "schema construction. A unified platform can reduce this load by "
        "presenting learners with consistent navigation, vocabulary, "
        "interaction patterns, and feedback mechanisms.",
    )
    add_paragraph(
        doc,
        "Motivationally, evidence indicates that young learners persist longer "
        "and engage more deeply when learning activities are gamified and "
        "supported by encouraging teacher relationships. Wallace and Leong "
        "(2020) and Roorda et al. (2017) demonstrate that motivation and "
        "engagement are closely linked to game-like activities and supportive "
        "teacher\u2013student interaction. Alotaibi (2024) and Zhonggen "
        "(2019) report moderate to large positive effects of game-based "
        "learning and serious games on engagement and emotional "
        "experience. Integrating gamified practice with lightweight teacher "
        "communication therefore directly addresses the motivational risks "
        "identified in Section 2.1.",
    )
    add_paragraph(
        doc,
        "Linguistically, research on technology-enhanced Mandarin "
        "instruction shows that well-designed online environments can "
        "improve targeted skills. Chen (2024) demonstrates that "
        "computer-aided visual feedback supports Mandarin tone production, "
        "and Wang (2024) shows that a theory-driven digital game can support "
        "Chinese character learning. Hongsamsibkao (2024) confirms that "
        "structured online Chinese courses improve listening and Pinyin "
        "outcomes. These findings collectively justify the inclusion of "
        "structured Pinyin, listening, and character activities in the proposed "
        "platform.",
    )
    add_paragraph(
        doc,
        "Technologically, recent literature on AI in education shows that "
        "bounded, supervised AI support can scaffold learners while "
        "preserving teacher authority. Du and Daniel (2024) report that "
        "AI-powered chatbots positively influence learner confidence and "
        "engagement when used as supplementary practice tools. Garz\u00f3n "
        "et al. (2025) and Wang et al. (2024) recommend explicit limits, "
        "transparency, and human oversight when AI is used with learners. "
        "Adopting an external AI API behind a backend proxy with teacher "
        "review for any publishable AI-generated content is therefore both "
        "technically realistic for an FYP and ethically appropriate.",
    )
    add_paragraph(
        doc,
        "Contextually, Malaysia's multilingual classroom culture and the "
        "presence of SJKC, KSSR POL Mandarin classes, and private tuition "
        "centres demand a system whose interface, content, and scaffolding "
        "respect local language realities. Rahmat et al. (2021) and Chua et al. "
        "(2024) emphasise that local Mandarin learners benefit from "
        "interfaces and instructional flows that acknowledge their other "
        "languages. Providing English, Bahasa Melayu, and Simplified "
        "Chinese interfaces, together with culturally responsive examples, "
        "directly addresses this contextual requirement.",
    )

    add_heading(doc, "2.3 Impact of the Problem", 2)
    add_paragraph(
        doc,
        "If the issues identified in Section 2.1 remain unaddressed, several "
        "negative consequences are likely to persist across the Malaysian "
        "Mandarin learning ecosystem.",
    )
    add_paragraph(
        doc,
        "For students, the continued use of fragmented, weakly scaffolded "
        "tools may reinforce low motivation, recurring tone and pronunciation "
        "errors, weak character recognition, and a tendency to memorise "
        "without meaning. Limited feedback can allow errors to fossilise, and "
        "fragmented digital experiences can cause attention to drift between "
        "platforms. Over time, learners may lose interest in Mandarin practice, "
        "weakening the very persistence that is essential for long-term "
        "language acquisition (Paee et al., 2024; Roorda et al., 2017).",
    )
    add_paragraph(
        doc,
        "For teachers, the absence of an integrated platform may continue to "
        "consume disproportionate time on routine tasks: preparing beginner "
        "materials, marking practice items, sending generic instructions, and "
        "tracking progress across multiple channels. This reduces the time "
        "teachers can devote to higher-value coaching, learner encouragement, "
        "and individualised feedback. Tuition centres in particular face high "
        "teacher turnover and limited preparation time, which compounds the "
        "issue.",
    )
    add_paragraph(
        doc,
        "For administrators of private classes, language centres, and "
        "after-school programmes, the lack of an integrated platform "
        "complicates content moderation, learner safety oversight, report "
        "handling, and analytics. Maintaining a safe and pedagogically "
        "appropriate environment for children becomes increasingly difficult "
        "as learner numbers grow without systematic tooling.",
    )
    add_paragraph(
        doc,
        "For parents and guardians, fragmented learning tools reduce "
        "visibility into what children are practising, which makes informed "
        "supervision and consent more difficult. Although parents are not a "
        "direct user role in the proposed prototype, child-appropriate design, "
        "consent disclosure, and child-safety pages remain essential to earn "
        "parental trust.",
    )
    add_paragraph(
        doc,
        "For the broader Malaysian educational technology ecosystem, the "
        "lack of integrated child-oriented prototypes means there are few "
        "practical reference designs that demonstrate how beginner Mandarin "
        "content, role-based educational workflows, multilingual access, and "
        "bounded AI support can be combined responsibly. The proposed "
        "study therefore has value not only as an academic FYP but also as a "
        "reference model that could inform future research and development "
        "in similar contexts within ASEAN.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 3 - Objectives
# ---------------------------------------------------------------------------

def write_objectives(doc: Document) -> None:
    add_heading(doc, "3.0 Objectives", 1)
    add_heading(doc, "3.1 General Objective", 2)
    add_paragraph(
        doc,
        "The general objective of this study is to design, develop, and "
        "evaluate a web-based prototype of an AI-Powered Interactive "
        "Chinese Learning Platform that supports beginner Mandarin learners "
        "aged 7 to 12 in Malaysia through structured content, child-friendly "
        "interactive activities, bounded AI assistance, teacher-managed "
        "workflows, and multilingual access within a coherent FYP-scale "
        "system.",
    )

    add_heading(doc, "3.2 Specific Objectives", 2)
    add_paragraph(
        doc,
        "Six SMART (Specific, Measurable, Achievable, Relevant, "
        "Time-bound) specific objectives operationalise the general objective. "
        "Each objective is mapped explicitly to a development phase and an "
        "evaluation source in Table 6.1.",
    )
    objectives = [
        ("SO1", "To analyse the instructional and usability needs of beginner "
                "Mandarin learners aged 7 to 12 in Malaysia, the workflow "
                "needs of Mandarin teachers, and the oversight needs of "
                "administrators by reviewing recent literature, comparing five "
                "existing learning platforms, and conducting at least one "
                "informal expert consultation by the end of Week 4."),
        ("SO2", "To design the proposed system architecture, database "
                "structure, role-based user interface flow, and core diagrams "
                "(use case, ER, component, deployment, activity, sequence, "
                "and state machine diagrams) for Student, Teacher, "
                "Administrator, and Guest users by the end of Week 7."),
        ("SO3", "To develop a full-stack web prototype using Vue 3, Spring "
                "Boot 3.5.x, and MySQL 8.4 LTS that implements secure JWT "
                "authentication, role-based access control, teacher "
                "verification workflow, learning content management, "
                "interactive learning activities, gamified rewards, and "
                "multilingual interfaces (English, Bahasa Melayu, Simplified "
                "Chinese) by the end of Week 18."),
        ("SO4", "To implement structured beginner Mandarin learning "
                "content and at least five interactive activity types\u2014"
                "flashcards, multiple-choice quizzes, matching games, "
                "listening exercises, and introductory tracing\u2014covering "
                "Pinyin, vocabulary, listening, simple sentence patterns, and "
                "introductory Chinese characters."),
        ("SO5", "To integrate bounded AI-assisted support through an "
                "external OpenAI-compatible API for at least seven "
                "supportive functions (chat tutor, vocabulary helper, sentence "
                "practice, beginner pronunciation guidance, beginner writing "
                "guidance, quiz or story generation, and teacher material "
                "drafting), with all publishable AI-generated content subject "
                "to teacher review."),
        ("SO6", "To evaluate the prototype on technical feasibility, usability "
                "(target System Usability Scale score \u2265 68), perceived "
                "learning support, and child-appropriate design through "
                "expert review (\u2265 1 Mandarin educator), task-based "
                "user testing (\u2265 5 supervised participants), and a "
                "Likert-scale questionnaire by the end of Week 22."),
    ]
    for code, text in objectives:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.left_indent = Cm(0.75)
        run = para.add_run(f"{code}. ")
        run.bold = True
        para.add_run(text)
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 4 - Scope of the Study
# ---------------------------------------------------------------------------

def write_scope(doc: Document) -> None:
    add_heading(doc, "4.0 Scope of the Study", 1)
    add_heading(doc, "4.1 System Scope", 2)
    add_paragraph(
        doc,
        "The proposed system is a web-based, full-stack software prototype "
        "developed for academic evaluation as a Final Year Project at "
        "Universiti Putra Malaysia. It is not intended for commercial release "
        "in its current form. The proposed system scope encompasses the "
        "following nine modules, each aligned with the specific objectives "
        "stated in Section 3.2.",
    )
    in_scope = [
        ("Authentication, Access Control, and Teacher Verification",
         "The system will support student and teacher self-registration, "
         "secure login, JWT-based session management, password reset, and "
         "role-based access control. Teacher accounts will go through a "
         "verification workflow before being permitted to publish or moderate "
         "learning content. Administrator accounts will be provisioned "
         "internally to maintain platform integrity."),
        ("Learning Content Management",
         "Verified teachers and administrators will be able to manage "
         "categories, levels, modules, lessons, activities, media resources, "
         "and prototype-level course or package structures. Uploads will be "
         "restricted by file type and size, and a content approval workflow "
         "will gate public visibility."),
        ("Interactive Learning Activities",
         "The system will include beginner-level Mandarin activities such as "
         "flashcards, multiple-choice quizzes, matching activities, listening "
         "exercises, and introductory tracing or short writing prompts. "
         "Activities will support practice in Pinyin, vocabulary, listening, "
         "simple sentence patterns, and introductory Chinese character "
         "recognition."),
        ("AI-Assisted Learning Features",
         "AI features will be integrated through a configurable external API "
         "(for example DeepSeek via an OpenAI-compatible interface). "
         "Planned AI-supported functions include a chat tutor, vocabulary "
         "helper, sentence practice support, beginner-level pronunciation "
         "guidance, beginner writing guidance, quiz or story generation, "
         "recommendation hints, learner progress insights, and teacher "
         "material drafting. AI is positioned as a supportive tool rather than a "
         "final authority on assessment."),
        ("Recorded and Live Learning Support",
         "Recorded learning will be delivered through uploaded or linked "
         "video and audio resources. Live learning will be coordinated "
         "through external meeting links such as Google Meet or Zoom; no "
         "internal video conferencing stack will be developed."),
        ("Progress Tracking and Reward Mechanisms",
         "The system will track progress through activity records, quiz "
         "attempts, mastery indicators, and progress summaries. Reward "
         "mechanisms will include points, badges, levels, daily streaks, and "
         "achievement tracking to encourage consistent practice."),
        ("Teacher Features",
         "Teacher functionality will include lesson and activity authoring, "
         "AI-assisted lesson drafting, student progress monitoring, structured "
         "feedback, and basic class activity coordination."),
        ("Administrator Features",
         "Administrator functionality will include user management, teacher "
         "verification, content moderation, report and review handling, "
         "support ticket management, audit logging, basic system "
         "configuration, and platform usage analytics."),
        ("Multilingual and Auxiliary Prototype Features",
         "The system will provide a multilingual interface in English, Bahasa "
         "Melayu, and Simplified Chinese. Auxiliary features may include "
         "in-app notifications, child-appropriate legal and safety pages, and "
         "sandbox payment workflows used solely for demonstration."),
    ]
    for idx, (title, body) in enumerate(in_scope, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)

    add_paragraph(
        doc,
        "The following items are explicitly outside the scope of this proposal "
        "to keep the project realistic for a Final Year Project timeline:",
    )
    out_of_scope = [
        ("Custom AI model training",
         "The proposed prototype will not train a proprietary model; all AI "
         "features will rely on configurable third-party APIs."),
        ("In-house video conferencing",
         "Live classes will use external meeting platforms; no internal "
         "video stack will be implemented."),
        ("Real financial transactions",
         "Cart, checkout, refund, subscription, and payout flows will be "
         "implemented as sandbox or simulated workflows only."),
        ("Production-level legal, security, and deployment hardening",
         "The prototype is intended for academic demonstration rather than "
         "commercial deployment, accessibility certification, or formal "
         "PDPA-compliance audit."),
        ("Native iOS and Android applications",
         "Only a responsive web application will be developed; native "
         "mobile builds are excluded."),
        ("Production-quality speech recognition and pronunciation scoring",
         "Pronunciation feedback will be implemented as a prototype "
         "experience; production-grade scoring engines such as those used "
         "in commercial platforms are out of scope."),
    ]
    for idx, (title, body) in enumerate(out_of_scope, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)

    add_heading(doc, "4.2 User Scope", 2)
    add_paragraph(
        doc,
        "The proposed system caters to four direct user categories and one "
        "indirect stakeholder group. Parents and guardians are treated as "
        "indirect supporters rather than as a separate platform role in this "
        "FYP-scale prototype.",
    )
    users = [
        ("Students",
         "Beginner Mandarin learners aged approximately 7 to 12. Students "
         "will use the platform for structured lessons, interactive activities, "
         "AI-assisted practice, progress tracking, and reward-based learning. "
         "Use is expected to occur under appropriate teacher or guardian "
         "supervision, particularly during evaluation."),
        ("Teachers",
         "Mandarin teachers, trainee teachers, or relevant educators who "
         "create and manage learning content, monitor student progress, "
         "provide feedback, and support learner engagement. Teachers must "
         "register and complete verification before publishing or moderating "
         "learning materials."),
        ("Administrators",
         "Platform administrators who handle user management, teacher "
         "verification, content moderation, report and support handling, audit "
         "logs, and platform-level analytics."),
        ("Guests",
         "Unauthenticated visitors who can access the public landing page, "
         "course discovery, FAQ, terms, privacy, and child-safety pages, but "
         "cannot access learning activities, AI features, private learning "
         "records, or role-protected functions."),
    ]
    for idx, (title, body) in enumerate(users, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "Parents and guardians are recognised as indirect supporters who "
        "may supervise children, provide consent where required, and review "
        "public guidance pages. A dedicated parent or guardian role is "
        "outside the scope of this prototype and is acknowledged as a "
        "potential future extension.",
    )

    add_heading(doc, "4.3 Dataset Scope", 2)
    add_paragraph(
        doc,
        "The dataset scope is defined narrowly to align with FYP "
        "requirements, the Personal Data Protection Act 2010 (Malaysia), "
        "and child-safety considerations. Five categories of data are in "
        "scope, and one category is explicitly excluded.",
    )
    datasets = [
        ("Instructional content dataset",
         "A bounded set of beginner Mandarin learning materials prepared "
         "or curated for demonstration purposes, including approximately "
         "200\u2013400 vocabulary items, 80\u2013150 Pinyin and tone "
         "items, 40\u201380 listening prompts, 60\u2013120 quiz items, "
         "20\u201340 introductory characters with stroke metadata, and "
         "demonstration sentence patterns. The dataset is intended for "
         "prototype demonstration and is not a comprehensive curriculum."),
        ("System operational dataset",
         "Account data (user identifiers, hashed credentials, role "
         "assignments), teacher verification records, lesson and activity "
         "definitions, quiz attempts, progress and reward records, "
         "notifications, support tickets, reports, reviews, and audit logs "
         "required for system operation."),
        ("Evaluation dataset",
         "Questionnaire responses, expert feedback notes, observation "
         "notes, task-completion timing, and anonymised usage logs "
         "collected during prototype testing solely for evaluation and design "
         "improvement. Personally identifying information will be minimised "
         "and stored in accordance with PDPA principles."),
        ("AI interaction dataset",
         "Limited learner and teacher prompts, system instructions, lesson "
         "context fragments, and AI-generated responses required to invoke "
         "external AI features. The proposed system will not train or fine-tune "
         "a proprietary AI model and will not retain prompts beyond the "
         "audit window required for system reliability."),
        ("Sandbox commerce dataset",
         "Mock payment records, simulated subscription states, and demo "
         "payout records used solely to demonstrate commerce workflows. "
         "No real card numbers, banking details, or customer financial "
         "information will be collected."),
    ]
    for idx, (title, body) in enumerate(datasets, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "Excluded data. The proposed prototype will not collect biometric "
        "data, real card or banking details, sensitive medical information, "
        "social media scraping data, or large-scale personal datasets "
        "unrelated to the prototype objectives. If young learners participate "
        "in evaluation, only the minimum data necessary for consented "
        "testing will be collected, and parental or guardian consent will be "
        "obtained where required by institutional ethics guidance.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 5 - Literature Review
# ---------------------------------------------------------------------------

def write_literature_review(doc: Document) -> None:
    add_heading(doc, "5.0 Literature Review", 1)
    add_heading(doc, "5.1 Introduction to Literature Review", 2)
    add_paragraph(
        doc,
        "This chapter reviews six interrelated bodies of literature that inform "
        "the design of the proposed platform: (i) beginner Mandarin learning "
        "challenges and motivation; (ii) online and technology-enhanced "
        "Chinese instruction; (iii) AI-assisted language learning; (iv) "
        "game-based learning, gamification, and serious games; (v) "
        "teacher\u2013student interaction and scaffolding; and (vi) cognitive "
        "and developmental considerations for young learners. The review then "
        "presents the theoretical framework that integrates these bodies of "
        "evidence (Section 5.3), critically compares five widely used existing "
        "Chinese learning systems against the requirements of the proposed "
        "platform (Section 5.4 and Table 5.1), articulates the research and "
        "design gaps the project intends to address (Table 5.2), and "
        "summarises the implications for system design (Section 5.5). Where "
        "possible, the review prioritises empirical and review studies "
        "published between 2021 and 2026 to ensure currency.",
    )

    add_heading(doc, "5.2 Related Works / Previous Studies", 2)

    add_heading(doc, "5.2.1 Mandarin Learning Challenges and Motivation", 3)
    add_paragraph(
        doc,
        "Mandarin presents distinct linguistic challenges for beginners "
        "because, unlike alphabetic systems, learners must concurrently "
        "master a Romanised phonetic system (Hanyu Pinyin), four lexical "
        "tones, morpho-syllabic characters, and vocabulary. Among Malaysian "
        "tertiary Mandarin learners, Paee et al. (2024) report demotivating "
        "factors including assessment pressure, content difficulty, "
        "perceived workload, anxiety, boredom, instructor\u2013peer dynamics, "
        "and the difficulty of Chinese characters. Although their participants "
        "were university students, the same drivers are likely to be "
        "consequential for younger beginners with weaker self-regulation. "
        "Wallace and Leong (2020) interviewed 23 primary EFL learners in "
        "Macau and found high motivation when learning featured games, "
        "songs, and encouraging teachers. Yang (2020) similarly identified "
        "curiosity, satisfaction, and stimulating activities as key motivators "
        "for Chinese online young language learners, with intrinsic motivation "
        "showing greater stability than extrinsic motivation. Together, these "
        "studies justify the integration of game-based activities and "
        "supportive teacher interaction in the proposed platform.",
    )

    add_heading(doc, "5.2.2 Online and Technology-Enhanced Chinese Instruction", 3)
    add_paragraph(
        doc,
        "Online and technology-enhanced Chinese instruction has expanded "
        "rapidly across the past decade. Liu (2018) describes a typology "
        "spanning online, classroom, and blended Chinese instruction, "
        "highlighting how learning management systems and e-communication "
        "tools blur the boundary between in-person and online teaching. Yang "
        "(2020) confirms rising parental investment in early online language "
        "learning in China, while observing limited scholarly attention to this "
        "young-learner segment. Hongsamsibkao (2024) developed a fully "
        "online Pinyin and listening course for Thai learners and reported "
        "positive outcomes for both achievement and satisfaction, suggesting "
        "that well-designed online instruction is feasible at the beginner "
        "level. Within Malaysia, Rahmat et al. (2021) surveyed 173 Mandarin "
        "learners and found that online instruction strengthened "
        "interactions among learners, instructor, and content while reducing "
        "the perceived Zone of Proximal Development gap. Chua et al. (2024) "
        "compared Malaysian and German Mandarin teaching and emphasised "
        "that technology integration is most effective when teacher flexibility "
        "and learner motivation are preserved. Fang et al. (2024) provide a "
        "systematic literature review of mobile-assisted Chinese learning "
        "between 2010 and 2022, observing benefits for speaking proficiency "
        "but a strong skew toward higher education samples. Collectively, "
        "these studies establish the feasibility of an online Mandarin platform "
        "while underscoring the need for designs tailored to younger Malaysian "
        "learners.",
    )

    add_heading(doc, "5.2.3 AI-Assisted Language Learning", 3)
    add_paragraph(
        doc,
        "Recent systematic reviews on AI in education describe rapid "
        "diffusion of AI applications for personalisation, formative feedback, "
        "and content generation, while flagging persistent concerns about "
        "ethics, transparency, teacher readiness, and overdependence "
        "(Garz\u00f3n et al., 2025; Wang et al., 2024). In language learning "
        "specifically, Du and Daniel (2024) synthesise evidence on "
        "AI-powered chatbots for English as a foreign language speaking "
        "practice and report positive effects on confidence, engagement, "
        "motivation, and selected pronunciation features. Although the "
        "evidence base is dominated by English contexts and adult learners, "
        "the underlying mechanisms\u2014repeated low-stakes practice, "
        "timely feedback, and personalised explanation\u2014are plausibly "
        "transferable to beginner Mandarin learners.",
    )
    add_paragraph(
        doc,
        "On pronunciation, Chen (2024) demonstrates that computer-aided "
        "visual and auditory feedback delivered through Praat improves "
        "Mandarin tone production among beginners, and Sun et al. (2023) "
        "report that automatic speech recognition combined with peer "
        "correction can support second-language pronunciation. Although the "
        "proposed prototype will not integrate research-grade automatic "
        "speech recognition, these studies illustrate that even simple "
        "AI-assisted feedback can yield benefits when carefully bounded. "
        "Accordingly, the proposed platform will treat AI features\u2014"
        "explanations, hints, sentence prompts, and beginner-level "
        "feedback\u2014as supplementary support rather than authoritative "
        "evaluation.",
    )

    add_heading(doc, "5.2.4 Game-Based Learning, Gamification, and Serious Games", 3)
    add_paragraph(
        doc,
        "Game-based and gamified strategies are well established in "
        "research with young learners. Alotaibi (2024) conducted a "
        "systematic review and meta-analysis of game-based learning in "
        "early childhood education and reported moderate to large effects "
        "on cognitive, motivational, and engagement outcomes. Zhonggen "
        "(2019) reviewed serious games research over a decade and found "
        "consistent positive effects on emotional state and persistence, "
        "while warning that misalignment between game mechanics and "
        "educational goals can dilute outcomes. Checa and Bustillo (2020) "
        "reviewed 135 immersive virtual reality serious games and concluded "
        "that interactivity, immersion, and learner-centred design are "
        "essential for success, though immersive VR itself is outside the "
        "scope of the proposed prototype. In the Chinese-character domain, "
        "Wang (2024) developed a theory-driven digital game and "
        "demonstrated how task design and engagement can support "
        "character acquisition. These findings collectively justify the "
        "incorporation of points, badges, levels, streaks, achievement "
        "tracking, matching games, and tracing-style activities into the "
        "proposed platform, while keeping each mechanic anchored to a "
        "specific educational goal.",
    )

    add_heading(doc, "5.2.5 Teacher\u2013Student Interaction and Scaffolding", 3)
    add_paragraph(
        doc,
        "Teacher involvement remains pivotal even in technology-rich "
        "learning environments. Roorda et al. (2017) provide a "
        "meta-analytic update covering 189 studies and approximately "
        "250,000 students, demonstrating that affective teacher\u2013student "
        "relationships are positively associated with engagement and "
        "achievement, with engagement mediating the relationship. This "
        "evidence is particularly relevant to primary-aged learners who "
        "depend more strongly on guidance and encouragement. Scaffolding "
        "theory provides a complementary lens: van de Pol et al. (2010) "
        "synthesise a decade of scaffolding research and identify "
        "contingency, fading, and transfer of responsibility as the three "
        "essential features of effective scaffolding. Vygotsky's (1978) Zone "
        "of Proximal Development underpins this work conceptually. The "
        "proposed platform will operationalise these insights through "
        "teacher-managed content workflows, AI hints that fade as learners "
        "progress, and teacher feedback channels that complement the "
        "automated activities.",
    )

    add_heading(doc, "5.2.6 Cognitive and Developmental Considerations for Young Learners", 3)
    add_paragraph(
        doc,
        "The cognitive and developmental profile of children aged 7 to 12 "
        "constrains what kinds of digital tasks are appropriate. Butler (2022) "
        "reviews young-learner cognitive development\u2014executive "
        "function, theory of mind, and metalinguistic awareness\u2014and "
        "their interaction with language learning, highlighting how digital "
        "technologies can both support and burden cognition. Sweller, van "
        "Merri\u00ebnboer, and Paas (2019) re-examine Cognitive Load "
        "Theory two decades after its formulation and emphasise that "
        "instructional design should respect the limits of working memory "
        "for novel material. Drawing on these insights, the proposed "
        "platform will employ short tasks, gradual progression from simple to "
        "complex, careful multimodal use (combining audio with visuals "
        "without redundant text), and concise, age-appropriate AI responses.",
    )

    add_heading(doc, "5.3 Theoretical Framework / Models", 2)
    add_paragraph(
        doc,
        "The proposed platform is grounded in four complementary "
        "theoretical perspectives. Together they guide design decisions "
        "across content sequencing, scaffolding, motivational mechanics, "
        "and AI bounding. Figure 5.1 presents the integrated conceptual "
        "framework that links beginner Mandarin learning challenges, "
        "platform components, and expected outcomes.",
    )
    add_figure_caption(
        doc,
        "Figure 5.1. Proposed conceptual framework linking beginner "
        "Mandarin learning challenges, platform components, and expected "
        "engagement and learning support outcomes.",
    )
    framework_items = [
        ("Cognitive Load Theory (Sweller et al., 2019).",
         "Activities are designed to limit extraneous cognitive load via "
         "short tasks, consistent navigation, modality-appropriate "
         "multimedia, and worked examples. For absolute beginners, the "
         "platform favours guided practice with hints over open-ended "
         "production tasks."),
        ("Zone of Proximal Development and Scaffolding (Vygotsky, 1978; "
         "van de Pol et al., 2010).",
         "Support is provided through teacher feedback, scaffolded "
         "activities, and AI hints calibrated to learner level. Scaffolding is "
         "expected to fade as learners gain confidence, transferring "
         "responsibility to the learner over time."),
        ("Motivation and Engagement among Young Language Learners "
         "(Roorda et al., 2017; Wallace & Leong, 2020; Alotaibi, 2024).",
         "Game-based activities, reward mechanisms, progress tracking, "
         "and supportive teacher interaction are integrated to encourage "
         "consistent practice and reduce demotivation, particularly for "
         "challenging skills such as tones and characters."),
        ("Bounded AI Support in Education (Du & Daniel, 2024; Garz\u00f3n "
         "et al., 2025; Wang et al., 2024).",
         "AI is treated as a bounded support mechanism rather than an "
         "autonomous teacher or assessment authority. Publishable "
         "AI-generated content is subject to teacher review, AI feedback is "
         "framed as supportive rather than evaluative, and access to AI is "
         "controlled via the backend to protect credentials and audit "
         "interactions."),
    ]
    for title, body in framework_items:
        para = doc.add_paragraph()
        run = para.add_run(title + " ")
        run.bold = True
        para.add_run(body)

    add_heading(doc, "5.4 Research Gap Analysis", 2)
    add_paragraph(
        doc,
        "This section analyses the gap addressed by the proposed platform "
        "in two stages. Section 5.4.1 critically compares existing systems "
        "from a functional perspective, while Section 5.4.2 maps the "
        "literature against the capabilities required by the proposed "
        "platform.",
    )

    add_heading(doc, "5.4.1 Comparison of Existing Systems", 3)
    add_paragraph(
        doc,
        "Five widely used Chinese-learning systems were reviewed because "
        "they represent dominant solution categories: gamified general "
        "language apps (Duolingo), Chinese-focused beginner platforms "
        "(HelloChinese, ChineseSkill), grammar-oriented language apps "
        "(LingoDeer), and Chinese-character handwriting tools (Skritter). "
        "Table 5.1 summarises their main strength, primary limitation in "
        "relation to the proposed platform, and relevance to the design "
        "decisions of this project. The comparison is functional rather than "
        "an empirical evaluation of pedagogical effectiveness.",
    )
    add_table_caption(
        doc,
        "Table 5.1. Critical comparison of selected existing Chinese learning "
        "systems.",
    )
    competitor_rows = [
        ["Duolingo",
         "Gamified bite-sized lessons with points, levels, streaks, and "
         "broad language coverage.",
         "Not tailored to Malaysian young learners; no teacher-managed "
         "workflow or admin oversight; AI features are general-purpose.",
         "Informs reward and streak design for the proposed platform."],
        ["HelloChinese",
         "Chinese-focused, with Pinyin practice, pronunciation drills, "
         "handwriting, video, and spaced repetition.",
         "Self-paced mobile app; lacks teacher verification, moderation, and "
         "Malaysian multilingual context.",
         "Informs Pinyin, handwriting, and listening activity design."],
        ["ChineseSkill",
         "Beginner-friendly Chinese mini-lessons covering Pinyin to grammar "
         "with games and quizzes.",
         "Individual learning app; limited teacher and admin integration; "
         "minimal Malaysian context.",
         "Informs quiz, matching-game, and listening-task design."],
        ["LingoDeer",
         "Grammar-oriented curriculum with native-speaker audio and clear "
         "explanations.",
         "Multi-language scope; not Malaysia-specific; lacks teacher "
         "verification and child-oriented design.",
         "Informs structured explanations and progression design."],
        ["Skritter",
         "Character-writing practice with handwriting recognition and "
         "stroke-level feedback.",
         "Focused narrowly on character writing; no integrated lesson "
         "system, AI tutor, or teacher workflow.",
         "Informs introductory tracing and writing practice."],
    ]
    add_table(
        doc,
        header=[
            "Existing system",
            "Main strength",
            "Limitation for this project",
            "Relevance to proposed platform",
        ],
        rows=competitor_rows,
        column_widths_cm=[3.0, 4.5, 4.5, 4.0],
    )

    add_heading(doc, "5.4.2 Mapping of Literature to Required Capabilities", 3)
    add_paragraph(
        doc,
        "The literature and competitor review reveal four interlocking gaps "
        "that the proposed platform is positioned to address.",
    )
    gaps = [
        ("Population gap.",
         "Most digital Chinese learning research targets higher education "
         "or general second-language audiences. Studies focused on "
         "beginner Mandarin learners aged 7 to 12 in Malaysia, with their "
         "specific multilingual context, remain comparatively scarce "
         "(Fang et al., 2024; Paee et al., 2024)."),
        ("Integration gap.",
         "Existing tools and studies typically focus on isolated skills\u2014"
         "vocabulary, pronunciation, characters, or chatbot interaction. "
         "Few platforms or studies integrate Pinyin, vocabulary, listening, "
         "characters, sentence creation, and gamification within a single "
         "beginner-oriented environment."),
        ("Pedagogical workflow gap.",
         "Many learner-facing language apps lack teacher-side tooling for "
         "authoring, AI-assisted drafting, progress monitoring, content "
         "moderation, and feedback. Yet such tooling is essential for "
         "young-learner contexts (Roorda et al., 2017; van de Pol et al., "
         "2010)."),
        ("Bounded AI design gap.",
         "AI in education research is expanding, but explicit, replicable "
         "design models for child-oriented language learning that bound AI "
         "behind teacher review remain limited (Garz\u00f3n et al., 2025; "
         "Wang et al., 2024)."),
    ]
    for title, body in gaps:
        para = doc.add_paragraph()
        run = para.add_run(title + " ")
        run.bold = True
        para.add_run(body)

    add_paragraph(
        doc,
        "Table 5.2 presents the research gap analysis matrix that compares "
        "existing solution categories with the capabilities required by the "
        "proposed platform. The comparison is descriptive and based on the "
        "literature reviewed in this chapter.",
    )
    add_table_caption(
        doc,
        "Table 5.2. Research gap analysis matrix between existing solution "
        "categories and the proposed platform.",
    )
    matrix_header = [
        "Solution category (sample sources)",
        "Beginner Mandarin",
        "Pinyin / tone",
        "Character",
        "AI tutoring",
        "Gamification",
        "Teacher tools",
        "Multilingual UI",
        "Child-appropriate",
    ]
    matrix_rows = [
        ["Generic LMS (Liu, 2018)",
         "\u2014", "\u2014", "\u2014", "\u2014", "\u2014",
         "\u2713", "Partial", "\u2014"],
        ["Mobile-assisted Chinese learning (Fang et al., 2024)",
         "Partial", "Partial", "\u2014", "\u2014", "Partial",
         "\u2014", "\u2014", "\u2014"],
        ["AI chatbots for EFL speaking (Du & Daniel, 2024)",
         "\u2014", "\u2014", "\u2014", "\u2713", "\u2014",
         "\u2014", "\u2014", "\u2014"],
        ["Computer-aided Pinyin / tone feedback (Chen, 2024; "
         "Hongsamsibkao, 2024)",
         "\u2713", "\u2713", "\u2014", "Partial", "\u2014",
         "\u2014", "\u2014", "\u2014"],
        ["Digital game for Chinese characters (Wang, 2024)",
         "\u2713", "\u2014", "\u2713", "\u2014", "\u2713",
         "\u2014", "\u2014", "Partial"],
        ["AI in education (Garz\u00f3n et al., 2025; Wang, S. et al., 2024)",
         "\u2014", "\u2014", "\u2014", "\u2713", "Partial",
         "Partial", "\u2014", "Partial"],
        ["Proposed platform (this project)",
         "\u2713", "\u2713", "\u2713", "\u2713", "\u2713",
         "\u2713", "\u2713", "\u2713"],
    ]
    add_table(
        doc,
        header=matrix_header,
        rows=matrix_rows,
        column_widths_cm=[4.6, 1.6, 1.5, 1.4, 1.5, 1.7, 1.6, 1.6, 1.7],
    )
    add_paragraph(
        doc,
        "Note. \u2713 = capability is a documented focus of the cited "
        "category; Partial = capability appears in some studies but is not "
        "central; \u2014 = capability is not a focus of the cited literature. "
        "The matrix summarises coverage rather than evaluating any specific "
        "product.",
        italic=True,
    )

    add_heading(doc, "5.5 Summary", 2)
    add_paragraph(
        doc,
        "The reviewed literature substantiates the design decisions of the "
        "proposed platform. Beginner Mandarin learning involves cognitive, "
        "motivational, pronunciation, character-related, and feedback-related "
        "challenges. Technology-enhanced instruction, AI-assisted support, "
        "game-based learning, teacher scaffolding, and cognitive load "
        "considerations each provide partial answers. However, no existing "
        "system or research artefact appears to combine these dimensions "
        "for young Malaysian learners within a teacher-managed, multilingual, "
        "and child-appropriate web prototype. The proposed platform is "
        "therefore positioned to address an integrated research and design "
        "gap by uniting structured beginner content, gamified practice, "
        "bounded AI support, teacher and administrator workflows, and "
        "multilingual access in one coherent FYP-scale prototype.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 6 - Methodology
# ---------------------------------------------------------------------------

def write_methodology(doc: Document) -> None:
    add_heading(doc, "6.0 Methodology", 1)
    add_heading(doc, "6.1 Introduction", 2)
    add_paragraph(
        doc,
        "This chapter outlines the methodology proposed for the design, "
        "development, and evaluation of the AI-Powered Interactive Chinese "
        "Learning Platform. It explains the overall research design, the "
        "development methodology, the data collection methods used for "
        "requirement gathering and prototype evaluation, the system design "
        "(architecture, tools, and technologies), the verification and "
        "validation strategy, and the proposed evaluation methods. Because "
        "this document is a proposal, all activities described in this chapter "
        "refer to planned rather than completed work.",
    )

    add_heading(doc, "6.2 Research Design", 2)
    add_paragraph(
        doc,
        "This study adopts Design and Development Research (DDR) as its "
        "overall research design. DDR is widely adopted in educational "
        "technology research because it focuses on the systematic design, "
        "development, and evaluation of an artefact that addresses an "
        "identified educational problem. In the present study, the artefact "
        "is the proposed web-based Mandarin learning prototype for young "
        "beginners in Malaysia, and the educational problem is the "
        "fragmented, weakly scaffolded, and limited-feedback nature of "
        "current beginner Mandarin learning support.",
    )
    add_paragraph(
        doc,
        "The DDR process is organised into four stages: (i) problem analysis "
        "and requirement identification, (ii) design, (iii) prototype "
        "development, and (iv) evaluation and refinement. Each stage maps "
        "onto specific objectives in Section 3.2, as summarised in Table 6.1 "
        "below.",
    )
    add_table_caption(
        doc,
        "Table 6.1. Mapping of specific objectives to research design stages "
        "and evaluation evidence.",
    )
    obj_stage_rows = [
        ["SO1", "Problem analysis and requirement identification",
         "Literature review notes, competitor matrix, expert consultation memo."],
        ["SO2", "Design", "Architecture diagram, ER diagram, use case diagram, "
                "wireframes, API outline."],
        ["SO3", "Prototype development", "Working frontend and backend, "
                "JWT authentication, RBAC enforcement, teacher verification "
                "workflow, multilingual interface."],
        ["SO4", "Prototype development", "Implemented activities (flashcards, "
                "quizzes, matching, listening, tracing) and content."],
        ["SO5", "Prototype development", "Integrated AI-assisted features and "
                "teacher review workflow."],
        ["SO6", "Evaluation and refinement", "Technical feasibility report, "
                "SUS scores, perceived learning support data, expert review "
                "notes."],
    ]
    add_table(
        doc,
        header=[
            "Specific objective",
            "DDR stage",
            "Primary evaluation evidence",
        ],
        rows=obj_stage_rows,
        column_widths_cm=[2.5, 5.5, 8.0],
    )

    add_heading(doc, "6.3 Development Methodology", 2)
    add_paragraph(
        doc,
        "Within the DDR research design, Agile incremental development is "
        "adopted as the implementation strategy. Agile is well suited to a "
        "Final Year Project because the proposed platform consists of "
        "several interrelated modules that can be designed, implemented, "
        "tested, and refined in manageable increments. Each increment "
        "includes a short planning cycle, implementation, automated and "
        "manual testing, supervisor review, and refinement.",
    )
    add_paragraph(
        doc,
        "The development plan is organised into six increments, each "
        "delivering an end-to-end vertical slice of functionality:",
    )
    increments = [
        ("Increment 1 \u2013 Foundation.",
         "Project setup, version control, build pipelines, environment "
         "configuration, database baseline, and routing structure."),
        ("Increment 2 \u2013 Security and role management.",
         "Student and teacher self-registration, login, JWT-based "
         "authentication, role-based access control, role-specific "
         "dashboards, and the teacher verification workflow."),
        ("Increment 3 \u2013 Learning content management.",
         "Categories, levels, modules, lessons, activities, media support, "
         "draft and approval workflow, and content moderation hooks."),
        ("Increment 4 \u2013 Interactive learning activities.",
         "Flashcards, multiple-choice quizzes, matching games, listening "
         "exercises, and introductory tracing."),
        ("Increment 5 \u2013 AI assistance, progress, and rewards.",
         "Bounded AI assistance through an external API, progress tracking, "
         "reward mechanisms, teacher-side AI drafting, and notifications."),
        ("Increment 6 \u2013 Administration, evaluation, and refinement.",
         "Moderation tooling, support tickets, multilingual polish, recorded "
         "and live learning support, sandbox commerce flows, analytics, "
         "verification, validation, documentation, and refinements."),
    ]
    for title, body in increments:
        para = doc.add_paragraph()
        run = para.add_run(title + " ")
        run.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "This staged approach ensures that essential learning experience "
        "and security workflows are built first, and that auxiliary features "
        "are added only after the core modules are stable. Figure 6.1 "
        "presents the overall flow of the proposed methodology.",
    )
    add_figure_caption(
        doc,
        "Figure 6.1. Proposed methodology combining Design and "
        "Development Research with Agile incremental development.",
    )

    add_heading(doc, "6.4 Data Collection Methods", 2)
    add_paragraph(
        doc,
        "Data collection serves two purposes: requirement gathering "
        "(informing system design) and prototype evaluation (informing "
        "refinement and academic assessment). Five complementary methods "
        "are proposed.",
    )
    methods = [
        ("Literature review",
         "Recent journal articles (prioritising 2021\u20132026), systematic "
         "reviews, and selected educational technology studies will be "
         "analysed to identify learner challenges, design principles, and "
         "system requirements relevant to beginner Mandarin learning. The "
         "literature review constitutes the primary requirement-gathering "
         "method for this proposal."),
        ("Informal expert consultation",
         "At least one Mandarin teacher or relevant educator in Malaysia "
         "will be approached for an informal consultation. Topics will "
         "include common beginner difficulties, classroom workflows, "
         "appropriate activities for the 7\u201312 age group, and any "
         "concerns about AI use. Consultation notes will be used solely for "
         "requirement and design improvement."),
        ("Likert-scale questionnaire",
         "A questionnaire informed by the System Usability Scale (SUS; "
         "Brooke, 1996; Lewis, 2018) will be used during prototype "
         "evaluation. Participants will rate ease of use, navigation clarity, "
         "interface suitability, AI helpfulness, motivational value of "
         "rewards, perceived learning support, and overall satisfaction on a "
         "five-point Likert scale. The questionnaire is provided in Appendix "
         "B."),
        ("System usage data",
         "During prototype testing, selected operational data will be "
         "captured for evaluation purposes only, including activity "
         "completion, quiz attempt records, AI request counts, reward "
         "events, and broad navigation patterns. Personally identifying "
         "information will be minimised."),
        ("Observation during testing",
         "With participant consent, the researcher will observe selected "
         "users completing representative tasks. Observation notes will "
         "capture navigation problems, task completion difficulties, "
         "confusion points, and visible usability barriers."),
    ]
    for idx, (title, body) in enumerate(methods, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "Where direct involvement of children aged 7 to 12 is feasible and "
        "approved, parental or guardian consent and supervision will be "
        "obtained according to UPM and KPM guidance. If direct child "
        "testing is not feasible within the FYP timeline, the study will use "
        "proxy users such as junior university students or older primary "
        "students, and will explicitly acknowledge this limitation in the "
        "final report. No unnecessary personal or sensitive data will be "
        "collected, in line with the Personal Data Protection Act 2010 "
        "(Malaysia).",
    )

    add_heading(doc, "6.5 System Design (Architecture, Tools, Technologies)", 2)

    add_heading(doc, "6.5.1 System Architecture", 3)
    add_paragraph(
        doc,
        "The proposed system adopts a three-tier web architecture "
        "comprising the presentation tier, the application tier, and the data "
        "tier. The architecture is summarised in Figure 6.2.",
    )
    add_figure_caption(
        doc,
        "Figure 6.2. Proposed three-tier system architecture with "
        "role-based access and external integrations.",
    )
    add_paragraph(
        doc,
        "Presentation tier. The presentation tier is delivered as a "
        "single-page application using Vue 3, Vite, and TypeScript. Element "
        "Plus provides accessible UI components, Pinia manages client-side "
        "state, Vue Router handles role-aware navigation, Axios mediates "
        "API communication, and Vue I18n provides multilingual interfaces "
        "in English, Bahasa Melayu, and Simplified Chinese. ECharts will "
        "render analytics dashboards.",
    )
    add_paragraph(
        doc,
        "Application tier. The application tier is implemented as a RESTful "
        "API using Spring Boot 3.5.x on Java 21. Spring Security and JWT "
        "provide stateless authentication and authorisation, Spring Data JPA "
        "manages persistence, and Bean Validation enforces request "
        "integrity. The backend acts as a controlled gateway for the external "
        "AI API, ensuring that API keys remain server-side and that all AI "
        "interactions can be audited. Provider clients for AI, payment, and "
        "email integrations are designed with safe fallbacks so that the "
        "system continues to function when external services are unavailable.",
    )
    add_paragraph(
        doc,
        "Data tier. MySQL 8.4 LTS serves as the primary relational database "
        "for accounts, role assignments, teacher verification records, lesson "
        "structures, activity records, progress and reward records, "
        "notifications, support data, and audit logs. To simplify local "
        "development, MySQL will be orchestrated through Docker Compose. "
        "Uploaded media will be stored in a backend-controlled directory "
        "with restricted access and size limits.",
    )

    add_heading(doc, "6.5.2 Main Modules", 3)
    add_paragraph(
        doc,
        "The proposed prototype is organised into eleven main modules: "
        "Authentication and Role-Based Access; Teacher Registration and "
        "Verification; Learning Content and Package Management; Student "
        "Learning Activities; AI-Assisted Learning Support; Progress "
        "Tracking and Reward Mechanisms; Teacher Feedback and "
        "Classroom Support; Administrator Management and Moderation; "
        "Notifications, Reports, Reviews, and Support Tickets; Multilingual "
        "Interface and Auxiliary Prototype Workflows; and Basic Analytics. "
        "Each module is mapped to one or more specific objectives "
        "(Section 3.2) and to one or more development increments "
        "(Section 6.3).",
    )

    add_heading(doc, "6.5.3 Tools and Technologies", 3)
    add_table_caption(doc, "Table 6.2. Proposed tools and technologies.")
    tools_rows = [
        ["Frontend framework", "Vue 3"],
        ["Build tool", "Vite"],
        ["Frontend language", "TypeScript"],
        ["UI component library", "Element Plus"],
        ["State management", "Pinia"],
        ["Routing", "Vue Router"],
        ["HTTP client", "Axios"],
        ["Internationalisation", "Vue I18n"],
        ["Charting", "ECharts"],
        ["Backend framework", "Spring Boot 3.5.x"],
        ["Backend language", "Java 21"],
        ["Security", "Spring Security with JWT"],
        ["Persistence", "Spring Data JPA"],
        ["Validation", "Jakarta Bean Validation"],
        ["Build tool (backend)", "Maven Wrapper"],
        ["Database", "MySQL 8.4 LTS"],
        ["AI integration",
         "Configurable external AI API (e.g., DeepSeek via OpenAI-compatible interface)"],
        ["Payment integration",
         "Stripe sandbox or simulated mode for prototype demonstration"],
        ["Email", "SMTP with safe no-op fallback when not configured"],
        ["Local infrastructure", "Docker Compose for MySQL"],
        ["Version control", "Git and GitHub"],
        ["Frontend testing", "Vitest"],
        ["Backend testing", "JUnit 5, Mockito, MockMvc"],
        ["Documentation", "Markdown and Microsoft Word"],
    ]
    add_table(
        doc,
        header=["Component", "Technology"],
        rows=tools_rows,
        column_widths_cm=[5.0, 11.0],
    )

    add_heading(doc, "6.5.4 Planned Diagrams", 3)
    add_paragraph(
        doc,
        "A complete set of diagrams will be prepared to support system "
        "analysis, system design, and implementation explanation. The "
        "planned diagrams include the conceptual framework diagram "
        "(Figure 5.1), the development methodology flow diagram (Figure "
        "6.1), the system architecture diagram (Figure 6.2), the system "
        "context diagram, the overall use case diagram, the entity "
        "relationship diagram, the component diagram, the deployment "
        "diagram, activity diagrams (student learning flow, teacher content "
        "creation flow, administrator teacher verification flow), sequence "
        "diagrams (user login, AI chat tutor request, quiz attempt, teacher "
        "content creation, administrator teacher verification), the state "
        "machine diagram for teacher verification status, the site map / "
        "navigation flow diagram, selected user interface wireframes, the "
        "Gantt chart, the Work Breakdown Structure, and a flowchart for "
        "quiz scoring and reward awarding. The list is reproduced in "
        "Appendix C.",
    )

    add_heading(doc, "6.6 Software Verification and Validation", 2)
    add_paragraph(
        doc,
        "Software verification and validation will be planned to ensure that "
        "the prototype is technically correct and appropriate for its "
        "intended users. Verification examines whether the system has been "
        "built correctly according to the specified requirements, while "
        "validation examines whether the system is suitable, usable, and "
        "useful within the FYP-scale beginner Mandarin learning context.",
    )

    add_heading(doc, "6.6.1 Verification Activities", 3)
    verification = [
        ("Backend unit testing",
         "Selected services and business logic will be tested with JUnit 5 "
         "and Mockito, prioritising authentication, role checks, and core "
         "domain rules."),
        ("Integration testing",
         "MockMvc-based integration tests will cover authentication, "
         "role-based access control, teacher verification, content "
         "workflows, quiz attempts, AI proxying, and progress updates."),
        ("API testing",
         "Backend endpoints will be tested under normal, error, and "
         "edge-case conditions to confirm consistent ApiResponse "
         "envelopes, validation handling, and authorisation results."),
        ("Frontend testing",
         "Vitest will be used for build checks, route guard checks, form "
         "validation checks, i18n key coverage, and selected component "
         "tests."),
        ("Database verification",
         "Schema migrations, foreign-key relationships, seed data, and "
         "data integrity will be verified against the planned ER diagram."),
        ("Prototype-level security checks",
         "Password hashing, JWT lifecycle, RBAC enforcement, "
         "server-side AI proxying, input validation, and exposure of "
         "sensitive fields will be reviewed."),
    ]
    for idx, (title, body) in enumerate(verification, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)

    add_heading(doc, "6.6.2 Validation Procedures", 3)
    validation = [
        ("User or proxy-user testing",
         "Selected participants will perform representative tasks on the "
         "prototype. Where children participate, parental or guardian "
         "consent and supervision will be obtained."),
        ("Teacher or educator review",
         "Feedback from at least one Mandarin teacher will be collected on "
         "content suitability, interaction flow, and child-friendly design."),
        ("Usability questionnaire",
         "A short Likert-scale questionnaire (informed by the SUS) will "
         "gather feedback on ease of use, navigation clarity, helpfulness, "
         "motivation, and suitability for young learners."),
        ("Observation notes",
         "Observation notes will record task difficulties, navigation "
         "issues, confusion points, and visible usability problems during "
         "testing."),
        ("Requirement checklist",
         "The completed prototype will be checked against the specific "
         "objectives and system scope to confirm that all planned "
         "functions have been addressed at prototype level."),
    ]
    for idx, (title, body) in enumerate(validation, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title}. ")
        run_title.bold = True
        para.add_run(body)

    add_heading(doc, "6.7 Evaluation Methods", 2)
    add_paragraph(
        doc,
        "Evaluation will examine the prototype along three dimensions: "
        "technical feasibility, usability, and perceived learning support. "
        "Each dimension is associated with explicit instruments and "
        "acceptance thresholds, summarised in Table 6.3.",
    )
    add_table_caption(
        doc,
        "Table 6.3. Evaluation criteria, instruments, and acceptance "
        "thresholds.",
    )
    eval_rows = [
        ["Technical feasibility",
         "JUnit/MockMvc backend tests, Vitest frontend tests, manual "
         "smoke tests, defect log.",
         "\u2265 95% pass rate across automated tests; all critical "
         "user journeys (registration, login, lesson play, AI tutor, "
         "quiz attempt, teacher publish, admin verify) demonstrably "
         "operational."],
        ["Usability",
         "Likert-scale questionnaire informed by the System Usability "
         "Scale (SUS), task completion observation.",
         "Mean SUS-equivalent score \u2265 68, which corresponds to the "
         "commonly reported benchmark for above-average usability "
         "(Lewis, 2018); \u2265 80% task completion rate."],
        ["Perceived learning support",
         "Likert-scale items, expert review, open-ended comments.",
         "Mean \u2265 4.0 on a 5-point Likert scale on perceived "
         "helpfulness for beginner Mandarin practice; positive expert "
         "review on age-appropriate design."],
        ["Child-appropriate design",
         "Expert review, observation, child-safety checklist.",
         "Expert confirmation that interface, tone, and content are "
         "suitable for children aged 7\u201312; no critical safety issues."],
    ]
    add_table(
        doc,
        header=["Evaluation dimension", "Instrument", "Acceptance threshold"],
        rows=eval_rows,
        column_widths_cm=[3.5, 5.5, 7.0],
    )
    add_paragraph(
        doc,
        "Open-ended comments from expert review and user testing will be "
        "analysed thematically at a descriptive level to identify recurring "
        "strengths, weaknesses, and improvement points. The evaluation "
        "does not aim to claim long-term Mandarin proficiency outcomes, "
        "but rather to assess whether the prototype is feasible, usable, and "
        "plausibly supportive for beginner Mandarin learning within the "
        "FYP scope.",
    )

    add_heading(doc, "6.8 Summary", 2)
    add_paragraph(
        doc,
        "In summary, the methodology integrates Design and Development "
        "Research with Agile incremental development to keep the project "
        "academically grounded yet practical for an FYP timeline. The "
        "proposed methods support requirement analysis, design, "
        "incremental development, verification, validation, and early-stage "
        "evaluation of the prototype, with explicit acceptance thresholds "
        "drawn from established educational technology measurement "
        "instruments such as the System Usability Scale.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 7 - ROI / ROV
# ---------------------------------------------------------------------------

def write_roi_rov(doc: Document) -> None:
    add_heading(doc, "7.0 Return on Investment / Return on Value", 1)
    add_heading(doc, "7.1 Cost-Benefit Analysis", 2)
    add_paragraph(
        doc,
        "Because this project is an academic Final Year Project prototype "
        "rather than a commercial product, its value is more appropriately "
        "framed as Return on Value (ROV) than as profit-driven Return on "
        "Investment (ROI). The primary objective is to generate educational, "
        "technical, and academic value rather than financial profit. "
        "Nevertheless, a cost-benefit analysis is presented below to "
        "demonstrate the project's economic feasibility within a typical "
        "Malaysian undergraduate budget. All amounts are in Malaysian "
        "Ringgit (MYR).",
    )
    add_table_caption(
        doc,
        "Table 7.1. Estimated cost structure for the proposed prototype "
        "(in MYR).",
    )
    cost_rows = [
        ["Development tools", "Visual Studio Code, IntelliJ IDEA Community, "
         "Git, GitHub, Java 21, Node.js, Maven, Vite",
         "Free", "All free or open-source."],
        ["Database", "MySQL 8.4 LTS via Docker Compose",
         "Free", "Local containerised setup."],
        ["UI library and frameworks", "Vue 3, Element Plus, Vue Router, "
         "Pinia, Axios, Vue I18n, ECharts; Spring Boot, Spring Security, "
         "Spring Data JPA",
         "Free", "Open-source ecosystem."],
        ["Hosting", "Localhost during development; optional free-tier "
         "deployment (e.g., Render, Railway) for demonstration",
         "0\u201360", "Free tier sufficient for demo."],
        ["External AI API", "DeepSeek (or compatible) usage during "
         "development and demonstration",
         "50\u2013120", "Pay-as-you-go; can be capped via configuration."],
        ["Stripe sandbox / Payment provider",
         "Test-mode credentials for sandbox demonstration",
         "Free", "No real payment processing."],
        ["Documentation and printing",
         "Final report printing and binding for submission",
         "20\u201340", "Hard copy as required by department."],
        ["Contingency", "Buffer for unforeseen development needs",
         "30\u201350", "Recommended reserve."],
        ["Estimated total (excluding student time)",
         "Aggregated lower- and upper-bound estimates from the rows above.",
         "100\u2013270",
         "Indicative range. Personal device, internet, and student time "
         "treated as sunk costs."],
    ]
    add_table(
        doc,
        header=["Cost item", "Description", "Estimated cost (MYR)", "Notes"],
        rows=cost_rows,
        column_widths_cm=[3.7, 5.8, 2.8, 3.7],
    )
    add_paragraph(
        doc,
        "The total out-of-pocket cost is therefore expected to remain "
        "modest, in the indicative range of MYR 100 to MYR 270, "
        "excluding the student's own device, existing internet access, and "
        "student time. The dominant project investment is the student's "
        "effort across approximately 24 weeks for requirement analysis, "
        "design, development, testing, evaluation, and documentation.",
    )

    add_heading(doc, "7.2 Expected Benefits (Tangible and Intangible)", 2)
    add_table_caption(
        doc,
        "Table 7.2. Expected tangible and intangible benefits of the "
        "proposed project.",
    )
    benefit_rows = [
        ["Tangible", "Working full-stack web prototype",
         "Demonstrable system covering authentication, RBAC, content "
         "management, AI assistance, gamified rewards, multilingual "
         "interface, and analytics."],
        ["Tangible", "Reusable software artefacts",
         "Database schema, API design, system diagrams, frontend "
         "components, and backend services that can inform follow-up "
         "research and development."],
        ["Tangible", "Documentation",
         "FYP proposal, final report, user guide, demonstration script, "
         "testing checklist, and proposal-style diagram pack."],
        ["Tangible", "Evaluation instruments",
         "Likert-scale usability questionnaire, observation notes template, "
         "expert interview guide, and requirement-to-feature traceability "
         "matrix."],
        ["Intangible", "Educational design knowledge",
         "Insights into how to translate beginner Mandarin learning needs "
         "into software requirements for young Malaysian learners."],
        ["Intangible", "AI-in-education insight",
         "A practical model for bounding AI assistance in a child-oriented "
         "context, contributing to ongoing discussions in the literature."],
        ["Intangible", "Software engineering competence",
         "Hands-on experience integrating modern frontend and backend "
         "stacks, secure authentication, multilingual UX, and external "
         "service integration."],
        ["Intangible", "Academic and career value",
         "A portfolio-grade FYP demonstrating ability to integrate research, "
         "design, and engineering for a non-trivial educational problem."],
    ]
    add_table(
        doc,
        header=["Type", "Benefit", "Description"],
        rows=benefit_rows,
        column_widths_cm=[2.0, 4.2, 9.8],
    )

    add_heading(doc, "7.3 Value Contribution to Stakeholders", 2)
    stakeholders = [
        ("Students.",
         "Beginner Mandarin learners may benefit from short, structured, "
         "and gamified activities; bounded AI hints; teacher feedback; "
         "progress visibility; and a child-appropriate multilingual "
         "interface that respects their existing language repertoire."),
        ("Teachers.",
         "Mandarin teachers in SJKC, KSSR POL classes, and tuition "
         "centres may benefit from AI-assisted material drafting, content "
         "management, progress monitoring, structured feedback channels, "
         "and a controlled publishing workflow that reduces routine "
         "preparation effort."),
        ("Administrators.",
         "Administrators of language centres, tuition operators, and "
         "after-school programmes may benefit from centralised tools for "
         "user management, teacher verification, content moderation, "
         "report and review handling, audit logging, and basic platform "
         "analytics."),
        ("Educational institutions and language centres.",
         "Public and private institutions may use the prototype as a "
         "reference design when considering how to integrate AI-assisted "
         "Mandarin learning, teacher workflows, and child-safety controls. "
         "The prototype does not claim institutional readiness or "
         "commercial deployability."),
        ("Parents and guardians.",
         "Although parents are not direct system users, they are indirect "
         "beneficiaries through child-appropriate design, transparent "
         "consent disclosures, and clearer supervision guidance."),
        ("Academic evaluators and the SEIS Department.",
         "Academic evaluators can assess the proposal and resulting "
         "prototype against software engineering, educational technology, "
         "and child-safety criteria. The project also contributes a local "
         "exemplar that may inform future SEIS FYP projects in "
         "education-related domains."),
        ("Universiti Putra Malaysia (UPM).",
         "The project enhances UPM's academic portfolio in educational "
         "technology research and demonstrates the integration of recent "
         "AI advances with culturally responsive design for the Malaysian "
         "multilingual context."),
    ]
    for idx, (title, body) in enumerate(stakeholders, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title} ")
        run_title.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "In summary, the proposed prototype is positioned to deliver value "
        "across multiple stakeholder groups by demonstrating an FYP-scale "
        "but research-informed approach to integrated, child-oriented, "
        "multilingual Mandarin learning support enhanced by bounded AI.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 8 - Project Planning
# ---------------------------------------------------------------------------

def write_project_planning(doc: Document) -> None:
    add_heading(doc, "8.0 Project Planning", 1)
    add_heading(doc, "8.1 Work Breakdown Structure", 2)
    add_paragraph(
        doc,
        "The Work Breakdown Structure (WBS) decomposes the project into "
        "five major work packages aligned with the DDR stages described "
        "in Section 6.2 and the Agile increments described in Section 6.3. "
        "Each WBS item maps to a methodology phase and a primary "
        "deliverable, as shown in Table 8.1.",
    )
    add_table_caption(
        doc,
        "Table 8.1. Proposed Work Breakdown Structure (WBS) mapped to "
        "the methodology.",
    )
    wbs_rows = [
        ["1.0", "Planning and requirement analysis",
         "Problem analysis", "Problem statement, objectives, scope."],
        ["1.1", "Literature review",
         "Problem analysis",
         "Literature review section and reference list."],
        ["1.2", "User and requirement identification",
         "Problem analysis",
         "Functional and non-functional requirements."],
        ["1.3", "Proposal preparation",
         "Problem analysis",
         "FYP proposal report submitted to the supervisor."],
        ["2.0", "System design", "Design",
         "Architecture, database, and UI design artefacts."],
        ["2.1", "Frontend design", "Design",
         "Vue page structure, wireframes, navigation flow."],
        ["2.2", "Backend design", "Design",
         "API and service-layer design, security plan."],
        ["2.3", "Database design", "Design",
         "ER diagram, table definitions, seed strategy."],
        ["3.0", "Core development",
         "Prototype development", "Working prototype across increments."],
        ["3.1", "Authentication and RBAC",
         "Prototype development",
         "Login, registration, JWT, role-based access."],
        ["3.2", "Role-based dashboards",
         "Prototype development",
         "Student, teacher, and admin dashboards."],
        ["3.3", "Learning content and packages",
         "Prototype development",
         "Categories, levels, modules, lessons, activities, packages."],
        ["3.4", "Interactive activities",
         "Prototype development",
         "Flashcards, quizzes, matching games, listening, tracing."],
        ["3.5", "AI-assisted support",
         "Prototype development",
         "AI tutor, vocabulary helper, sentence practice, beginner "
         "feedback, teacher drafting."],
        ["3.6", "Progress and rewards",
         "Prototype development",
         "Progress tracking, points, badges, levels, streaks."],
        ["3.7", "Notifications, support, moderation",
         "Prototype development",
         "Notifications, support tickets, moderation actions, audit logs."],
        ["3.8", "Multilingual, live/recorded, and auxiliary modules",
         "Prototype development",
         "i18n, recorded learning, live links, sandbox payment, legal "
         "pages, analytics dashboards."],
        ["4.0", "Verification and validation",
         "Evaluation and refinement",
         "Test reports and bug fixes."],
        ["4.1", "Backend testing",
         "Evaluation and refinement",
         "Unit and integration test outputs."],
        ["4.2", "Frontend testing",
         "Evaluation and refinement",
         "Build and component test outputs."],
        ["4.3", "User testing and evaluation",
         "Evaluation and refinement",
         "Usability questionnaire, expert review, observation notes."],
        ["5.0", "Documentation and finalisation",
         "Evaluation and refinement",
         "Final report, user guide, demonstration script."],
        ["5.1", "Final report writing",
         "Evaluation and refinement",
         "FYP final report."],
        ["5.2", "Presentation preparation",
         "Evaluation and refinement",
         "Slides, demonstration video, rehearsal notes."],
    ]
    add_table(
        doc,
        header=["WBS", "Task", "Methodology phase", "Deliverable"],
        rows=wbs_rows,
        column_widths_cm=[1.4, 4.4, 3.4, 6.8],
    )

    add_paragraph(
        doc,
        "A risk register accompanies the WBS to anticipate and mitigate "
        "common pitfalls in educational technology FYPs.",
    )
    add_table_caption(doc, "Table 8.2. Risk register and mitigation plan.")
    risk_rows = [
        ["R1", "Scope creep due to broad feature set",
         "High",
         "Adopt the six-increment plan; defer auxiliary features until core "
         "modules are stable; align scope with supervisor weekly."],
        ["R2", "External AI API outage or rate limits",
         "Medium",
         "Implement safe fallback responses; cap AI request volumes; "
         "make AI features optional during demos."],
        ["R3", "Limited access to children for usability testing",
         "Medium",
         "Engage Mandarin teachers and proxy users; secure parental "
         "consent in advance; document the limitation transparently."],
        ["R4", "Data privacy concerns under PDPA when handling child data",
         "Medium",
         "Minimise data collection; anonymise records; obtain consent; "
         "store data on local development environment only."],
        ["R5", "Integration complexity (frontend, backend, AI, payment, SMTP)",
         "Medium",
         "Use provider client interfaces with safe fallbacks; cover critical "
         "flows with automated tests; prioritise core flows first."],
        ["R6", "Time overrun against the 24-week schedule",
         "Medium",
         "Hold weekly supervisor check-ins; use Agile sprints; maintain a "
         "buffer of two weeks for evaluation and final polish."],
    ]
    add_table(
        doc,
        header=["ID", "Risk", "Likelihood / Impact", "Mitigation"],
        rows=risk_rows,
        column_widths_cm=[1.0, 5.5, 2.5, 7.0],
    )

    add_heading(doc, "8.2 Gantt Chart", 2)
    add_paragraph(
        doc,
        "The project is planned over approximately 24 weeks across "
        "Semester 2, Session 2025/2026. The early phase focuses on "
        "planning, problem analysis, literature review, and proposal "
        "preparation. The middle phase focuses on incremental prototype "
        "development. The later phase focuses on verification, validation, "
        "evaluation, and final reporting. Figure 8.1 presents the proposed "
        "Gantt chart, and Appendix A provides the detailed weekly "
        "allocation. The schedule remains indicative and may be adjusted "
        "according to supervisor feedback, evaluation findings, and actual "
        "development progress.",
    )
    add_figure_caption(
        doc,
        "Figure 8.1. Proposed Final Year Project Gantt chart spanning 24 "
        "weeks across the semester.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 9 - Conclusion
# ---------------------------------------------------------------------------

def write_conclusion(doc: Document) -> None:
    add_heading(doc, "9.0 Conclusion", 1)
    add_heading(doc, "9.1 Summary of Proposal", 2)
    add_paragraph(
        doc,
        "This proposal outlines the design, development, and evaluation of "
        "an Artificial Intelligence (AI)-Powered Interactive Chinese "
        "Learning Platform for young beginners aged 7 to 12 in Malaysia. "
        "The proposed project is positioned as an FYP-scale, web-based "
        "prototype that responds to five interrelated problems documented "
        "in the literature: cognitive overload from simultaneous Mandarin "
        "skill demands; limited formative feedback; fragmentation of "
        "digital learning tools; heavy teacher workload for beginner "
        "support; and the risks of using general-purpose AI with young "
        "learners. The proposal motivates the work pedagogically, "
        "motivationally, technologically, and personally, and frames it "
        "within UPM's Bachelor of Software Engineering with Honours "
        "programme.",
    )
    add_paragraph(
        doc,
        "To address these problems, the proposed prototype will integrate "
        "structured beginner Mandarin content, gamified interactive "
        "activities, bounded AI assistance, multilingual interfaces, "
        "teacher-managed content workflows, and basic administrator "
        "moderation and analytics into one coherent web application. The "
        "system will deliberately exclude custom AI model training, in-house "
        "video conferencing, real financial transactions, native mobile apps, "
        "and production-grade speech-recognition scoring, in keeping with "
        "FYP scope.",
    )
    add_paragraph(
        doc,
        "The methodology combines Design and Development Research "
        "with Agile incremental development. Six development increments "
        "deliver an end-to-end vertical slice of functionality at each stage, "
        "while DDR provides the overall research framing. Evaluation will "
        "examine technical feasibility (\u2265 95% automated test pass "
        "rate; all critical journeys operational), usability (SUS "
        "score \u2265 68; \u2265 80% task completion rate), perceived "
        "learning support (mean \u2265 4.0 on a 5-point Likert scale), and "
        "child-appropriate design (positive expert review and no critical "
        "safety issues). The project is planned over 24 weeks, with a risk "
        "register and mitigation plan in place.",
    )
    add_paragraph(
        doc,
        "Overall, the proposal demonstrates how software engineering "
        "practices, educational technology research, and bounded AI "
        "ethics can be integrated into a coherent and academically "
        "defensible FYP for the Malaysian Mandarin-learning context.",
    )

    add_heading(doc, "9.2 Expected Contributions", 2)
    contributions = [
        ("Academic contribution.",
         "The project contributes an FYP-level study that connects "
         "research on Mandarin learning challenges, cognitive load, "
         "scaffolding, learner motivation, AI-assisted learning, and "
         "gamification with the design of a child-friendly learning prototype "
         "for the Malaysian multilingual context."),
        ("Technical contribution.",
         "The project will produce a full-stack prototype using Vue 3, "
         "Spring Boot 3.5.x, MySQL 8.4 LTS, JWT-based security, RBAC, "
         "and external AI API integration. The system design, database "
         "structure, API structure, and diagrams may serve as reusable "
         "technical references for similar educational technology projects."),
        ("Design contribution.",
         "The project contributes a coherent design model that combines "
         "beginner Mandarin content, interactive activities, reward "
         "mechanisms, multilingual interfaces, teacher involvement, "
         "administrator moderation, and bounded AI support. This model "
         "may inform similar projects within ASEAN."),
        ("Ethical and policy contribution.",
         "By bounding AI behind teacher review, minimising data "
         "collection, requiring parental consent for child testing, and "
         "explicitly excluding production-grade financial and "
         "speech-recognition functionality, the proposal contributes a "
         "responsible-by-design template for child-oriented AI in "
         "education within the Malaysian regulatory environment "
         "(including the Personal Data Protection Act 2010)."),
        ("Practical prototype contribution.",
         "The project will deliver a working prototype that demonstrates "
         "how young learners, teachers, and administrators could interact "
         "within a controlled Mandarin learning environment. It also "
         "delivers planned evaluation instruments such as interview "
         "questions, usability questionnaire items, observation notes, "
         "and requirement checklists."),
    ]
    for idx, (title, body) in enumerate(contributions, start=1):
        para = doc.add_paragraph()
        run_title = para.add_run(f"{idx}. {title} ")
        run_title.bold = True
        para.add_run(body)
    add_paragraph(
        doc,
        "In summary, the project is expected to contribute a realistic, "
        "research-informed, and technically feasible prototype design and "
        "evaluation plan rather than a production-ready product or a claim "
        "of proven long-term Mandarin learning improvement.",
    )
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 10 - References
# ---------------------------------------------------------------------------

def write_references(doc: Document) -> None:
    add_heading(doc, "10.0 References", 1)
    add_paragraph(
        doc,
        "The reference list follows the Publication Manual of the American "
        "Psychological Association, 7th edition (American Psychological "
        "Association, 2020). The list prioritises empirical and review "
        "studies published between 2021 and 2026. Foundational older "
        "citations are retained where they remain canonical for the "
        "theoretical framework. Competitor system websites are listed "
        "separately as informational resources rather than research "
        "evidence.",
        italic=True,
    )

    references = [
        "Alotaibi, M. S. (2024). Game-based learning in early childhood "
        "education: A systematic review and meta-analysis. Frontiers in "
        "Psychology, 15, 1307881. "
        "https://doi.org/10.3389/fpsyg.2024.1307881",

        "American Psychological Association. (2020). Publication manual of "
        "the American Psychological Association (7th ed.). "
        "https://doi.org/10.1037/0000165-000",

        "Brooke, J. (1996). SUS: A \u201cquick and dirty\u201d usability "
        "scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, & I. L. "
        "McClelland (Eds.), Usability evaluation in industry (pp. 189\u2013"
        "194). Taylor & Francis.",

        "Butler, Y. G. (2022). Cognition and young learners' language "
        "development. In M. Schwartz (Ed.), Handbook of early language "
        "education (pp. 29\u201355). Springer. "
        "https://doi.org/10.1007/978-3-030-91662-6_2",

        "Checa, D., & Bustillo, A. (2020). A review of immersive virtual "
        "reality serious games to enhance learning and training. Multimedia "
        "Tools and Applications, 79(9\u201310), 5501\u20135527. "
        "https://doi.org/10.1007/s11042-019-08348-9",

        "Chen, M. (2024). Computer-aided feedback on the pronunciation of "
        "Mandarin Chinese tones: Using Praat to promote multimedia foreign "
        "language learning. Computer Assisted Language Learning, 37(1\u20132), "
        "363\u2013388. "
        "https://doi.org/10.1080/09588221.2022.2037652",

        "Chua, N. A., Ahmad, M. B., Ismail, I. L. M., Abdul Mutalib, N. B., "
        "Saputra, J., Goh, Y. S., & Ngo, K. L. (2024). Technology integration "
        "in learning Mandarin Chinese as a foreign language in Malaysia and "
        "Germany. International Journal of Advanced and Applied Sciences, "
        "11(9), 192\u2013202. "
        "https://doi.org/10.21833/ijaas.2024.09.021",

        "Du, J., & Daniel, B. K. (2024). Transforming language education: A "
        "systematic review of AI-powered chatbots for English as a foreign "
        "language speaking practice. Computers and Education: Artificial "
        "Intelligence, 6, 100230. "
        "https://doi.org/10.1016/j.caeai.2024.100230",

        "Fang, J., Chew, F. P., & Shaharom, M. S. N. (2024). Mobile-assisted "
        "learning of Chinese as a second/foreign language abroad: A "
        "systematic literature review of studies between 2010 and 2022. "
        "Knowledge Management & E-Learning, 16(3), 501\u2013520. "
        "https://doi.org/10.34105/j.kmel.2024.16.023",

        "Garz\u00f3n, J., Pati\u00f1o, E., & Marulanda, C. (2025). Systematic "
        "review of artificial intelligence in education: Trends, benefits, and "
        "challenges. Multimodal Technologies and Interaction, 9(8), 84. "
        "https://doi.org/10.3390/mti9080084",

        "Hongsamsibkao, N. (2024). Development of listening comprehension "
        "and pronunciation of Pinyin characters in Mandarin Chinese through "
        "online language teaching: A study on Chinese language students at "
        "the Faculty of Education, Uttaradit Rajabhat University. Journal of "
        "Education and Learning, 13(4), 194\u2013200. "
        "https://doi.org/10.5539/jel.v13n4p194",

        "Kementerian Pendidikan Malaysia. (2017). Kurikulum Standard "
        "Sekolah Rendah (KSSR Semakan 2017): Bahasa Cina Sekolah "
        "Kebangsaan dan Sekolah Jenis Kebangsaan Cina. Bahagian "
        "Pembangunan Kurikulum.",

        "Lewis, J. R. (2018). The System Usability Scale: Past, present, and "
        "future. International Journal of Human\u2013Computer Interaction, "
        "34(7), 577\u2013590. "
        "https://doi.org/10.1080/10447318.2018.1455307",

        "Liu, S. (2018). Teaching and learning Chinese language online: "
        "What and why? International Chinese Language Education, 3(2), "
        "11\u201322.",

        "Paee, R., Din, W. A., & Zainuddin, A. (2024). Demotivation in "
        "Mandarin as a foreign language among Malaysian tertiary students. "
        "Journal of Nusantara Studies, 9(2), 40\u201356. "
        "https://doi.org/10.24200/jonus.vol9iss2pp40-56",

        "Rahmat, N. H., Mok, S. S., Lau, S. K., & Ling, T. S. (2021). An "
        "investigation of how online learning reduces ZPD in Mandarin "
        "language classrooms. International Journal of Education, 13(1), "
        "1\u201317. https://doi.org/10.5296/ije.v13i1.18399",

        "Roorda, D. L., Jak, S., Zee, M., Oort, F. J., & Koomen, H. M. Y. "
        "(2017). Affective teacher\u2013student relationships and students' "
        "engagement and achievement: A meta-analytic update and test of "
        "the mediating role of engagement. School Psychology Review, "
        "46(3), 239\u2013261. "
        "https://doi.org/10.17105/SPR-2017-0035.V46-3",

        "Sun, W., Shadiev, R., & Li, Y. (2023). The impact of automatic "
        "speech recognition technology on second language pronunciation "
        "and speaking skills of EFL learners: A mixed methods investigation. "
        "Frontiers in Psychology, 14, 1210187. "
        "https://doi.org/10.3389/fpsyg.2023.1210187",

        "Sweller, J., van Merri\u00ebnboer, J. J. G., & Paas, F. (2019). "
        "Cognitive architecture and instructional design: 20 years later. "
        "Educational Psychology Review, 31(2), 261\u2013292. "
        "https://doi.org/10.1007/s10648-019-09465-5",

        "van de Pol, J., Volman, M., & Beishuizen, J. (2010). Scaffolding in "
        "teacher\u2013student interaction: A decade of research. Educational "
        "Psychology Review, 22(3), 271\u2013296. "
        "https://doi.org/10.1007/s10648-010-9127-6",

        "Vygotsky, L. S. (1978). Mind in society: The development of higher "
        "psychological processes (M. Cole, V. John-Steiner, S. Scribner, & "
        "E. Souberman, Eds.). Harvard University Press.",

        "Wallace, M. P., & Leong, E. I. L. (2020). Exploring language "
        "learning motivation among primary EFL learners. Journal of "
        "Language Teaching and Research, 11(2), 221\u2013230. "
        "https://doi.org/10.17507/jltr.1102.10",

        "Wang, S., Wang, F., Zhu, Z., Wang, J., Tran, T., & Du, Z. (2024). "
        "Artificial intelligence in education: A systematic literature review. "
        "Expert Systems with Applications, 252, 124167. "
        "https://doi.org/10.1016/j.eswa.2024.124167",

        "Wang, T. (2024). Designing a digital game for Chinese character "
        "learning: A theory-driven practice approach. Education Sciences, "
        "14(12), 1366. https://doi.org/10.3390/educsci14121366",

        "Yang, J. (2020). Learning status, trends and issues of online young "
        "language learner in China. Frontiers in Educational Research, 3(1), "
        "14\u201320. https://doi.org/10.25236/FER.2020.030103",

        "Zhonggen, Y. (2019). A meta-analysis of use of serious games in "
        "education over a decade. International Journal of Computer Games "
        "Technology, 2019, 4797032. https://doi.org/10.1155/2019/4797032",
    ]
    for entry in references:
        add_reference(doc, entry)

    add_paragraph(doc, "")
    add_plain_paragraph(doc, "Competitor system websites (informational only):", bold=True)
    competitor_refs = [
        "ChineseSkill. (n.d.). ChineseSkill: Learn Chinese Mandarin apps. "
        "Retrieved May 13, 2026, from https://www.chineseskill.com/",
        "Duolingo. (n.d.). The world's most popular way to learn. Retrieved "
        "May 13, 2026, from https://www.duolingo.com/",
        "HelloChinese. (n.d.). HelloChinese - No. 1 Chinese learning app. "
        "Retrieved May 13, 2026, from https://www.hellochinese.cc/",
        "LingoDeer. (n.d.). Learn Japanese, Korean, Chinese, and more. "
        "Retrieved May 13, 2026, from https://www.lingodeer.com/",
        "Skritter. (n.d.). Skritter - Learn to write Chinese and Japanese "
        "characters. Retrieved May 13, 2026, from https://skritter.com/",
    ]
    for entry in competitor_refs:
        add_reference(doc, entry)
    add_page_break(doc)


# ---------------------------------------------------------------------------
# Section 11 - Appendices
# ---------------------------------------------------------------------------

def write_appendices(doc: Document) -> None:
    add_heading(doc, "11.0 Appendices", 1)

    add_heading(doc, "Appendix A. Project Timeline Details", 2)
    add_paragraph(
        doc,
        "The proposed project spans approximately 24 weeks. The "
        "indicative weekly allocation is presented below. The schedule is "
        "indicative and may be adjusted based on supervisor feedback, "
        "evaluation outcomes, and project realities.",
    )
    timeline_rows = [
        ["1\u20132", "Project initiation",
         "Background reading, supervisor briefing, scope confirmation."],
        ["3\u20134", "Problem analysis and literature review",
         "Refined problem statement, draft literature review, "
         "competitor analysis."],
        ["5\u20136", "Requirement and design",
         "Functional and non-functional requirements, use cases, "
         "wireframes."],
        ["7", "Architecture and database design",
         "Architecture diagram, ER diagram, API outline, environment "
         "setup."],
        ["8\u20139", "Authentication and RBAC (Increment 2)",
         "Registration, login, JWT, role-based access, dashboards."],
        ["10\u201311", "Teacher verification and content management "
         "(Increment 3)",
         "Verification workflow, modules, lessons, activities, packages."],
        ["12\u201313", "Interactive activities (Increment 4)",
         "Flashcards, quizzes, matching, listening, tracing."],
        ["14\u201315", "AI-assisted features and rewards (Increment 5)",
         "AI tutor, vocabulary helper, sentence practice, beginner "
         "feedback, points/badges."],
        ["16\u201317", "Notifications, support, moderation (Increment 5/6)",
         "Notifications, support tickets, moderation actions, audit logs."],
        ["18\u201319", "Multilingual, live/recorded, sandbox commerce, "
         "analytics (Increment 6)",
         "i18n, recorded learning, live links, sandbox payment, "
         "analytics dashboards."],
        ["20", "Polish and accessibility checks",
         "UX refinement, accessibility heuristics, content review."],
        ["21", "Verification and validation",
         "Backend tests, frontend tests, manual smoke tests, defect "
         "fixes."],
        ["22", "User testing and evaluation",
         "Expert review, supervised user testing, questionnaire, "
         "observation notes."],
        ["23", "Final report writing",
         "Final report draft, evaluation analysis, contributions write-up."],
        ["24", "Presentation and submission",
         "Slides, demonstration video, rehearsal, submission."],
    ]
    add_table(
        doc,
        header=["Week(s)", "Phase / Increment", "Key activities and outputs"],
        rows=timeline_rows,
        column_widths_cm=[1.8, 5.5, 8.7],
    )

    add_heading(doc, "Appendix B. Survey / Interview Instruments", 2)

    add_heading(doc,
                "B.1 Sample Interview Questions for Teacher / Education "
                "Expert", 3)
    interview_questions = [
        "From your experience teaching Mandarin in Malaysia, what are the "
        "main difficulties faced by beginner learners aged 7 to 12?",
        "Which area is usually the hardest for beginners\u2014Pinyin, tones, "
        "vocabulary, sentence patterns, or Chinese characters\u2014and why?",
        "What pronunciation or tone errors do beginner learners commonly "
        "make in the Malaysian classroom?",
        "Which teaching strategies tend to be most effective for children "
        "learning Mandarin in SJKC, KSSR POL classes, or tuition centres?",
        "In your view, can a structured online platform meaningfully "
        "support beginner Mandarin learners? Why or why not?",
        "Which features should a beginner-friendly Mandarin platform for "
        "young Malaysian learners include?",
        "How important is teacher feedback in beginner Mandarin learning, "
        "and how could a digital tool support that?",
        "Do you think AI-assisted features (chat tutor, vocabulary helper, "
        "pronunciation guidance) could be useful for young learners? What "
        "are your concerns?",
        "What role can gamified elements (points, badges, daily streaks) "
        "play in maintaining young learners' motivation?",
        "Are there any safety, content moderation, or child-appropriate "
        "design considerations you would highlight to a developer "
        "building such a platform?",
        "What multilingual considerations are important in a Malaysian "
        "context (English, Bahasa Melayu, Mandarin)?",
        "How might administrators of language centres or tuition operators "
        "benefit from an integrated platform?",
    ]
    for question in interview_questions:
        add_number(doc, question)

    add_heading(doc, "B.2 Sample Likert-Scale Usability Questionnaire", 3)
    add_paragraph(
        doc,
        "For each statement, the respondent indicates their level of "
        "agreement on a five-point scale: 1 = Strongly Disagree, "
        "2 = Disagree, 3 = Neutral, 4 = Agree, 5 = Strongly Agree.",
    )
    questionnaire_items = [
        "The platform was easy to use.",
        "The navigation was clear.",
        "The interface looked suitable for young learners.",
        "The AI chat tutor was helpful for practising Mandarin.",
        "The vocabulary helper was useful.",
        "The quizzes were easy to understand.",
        "The matching games made vocabulary practice more interesting.",
        "The listening activities helped me practise Mandarin listening.",
        "The tracing activity helped me practise Chinese characters.",
        "The progress tracking helped me see what I had completed.",
        "The points and badges encouraged me to continue practising.",
        "The daily streak motivated me to practise more often.",
        "The interface in English / Bahasa Melayu / Simplified Chinese was "
        "clear.",
        "I felt safe and comfortable while using the platform.",
        "I would like to use this platform again in the future.",
        "Overall, I am satisfied with the platform.",
    ]
    for item in questionnaire_items:
        add_number(doc, item)
    add_paragraph(
        doc,
        "Open-ended follow-up prompts: \"What did you like most?\", "
        "\"What did you find confusing?\", and \"What would you like to "
        "see improved?\".",
    )

    add_heading(doc, "B.3 SUS-Inspired Items (10 statements)", 3)
    sus_items = [
        "I think that I would like to use this platform frequently.",
        "I found the platform unnecessarily complex.",
        "I thought the platform was easy to use.",
        "I think that I would need the support of a teacher to use this "
        "platform.",
        "I found the various functions in this platform were well integrated.",
        "I thought there was too much inconsistency in this platform.",
        "I would imagine that most children would learn to use this "
        "platform very quickly.",
        "I found the platform very cumbersome to use.",
        "I felt very confident using the platform.",
        "I needed to learn a lot of things before I could get going with this "
        "platform.",
    ]
    for item in sus_items:
        add_number(doc, item)
    add_paragraph(
        doc,
        "Items adapted from Brooke (1996) and Lewis (2018). For child "
        "respondents, language will be simplified, and items may be "
        "answered with parental or teacher assistance.",
        italic=True,
    )

    add_heading(doc, "Appendix C. Additional Diagrams", 2)
    add_paragraph(
        doc,
        "The following diagrams will be produced and refined during the "
        "system analysis and design phases. All diagrams will be presented "
        "as proposal-style artefacts in the final report:",
    )
    diagrams = [
        "Conceptual Framework Diagram (Figure 5.1).",
        "Development Methodology Flow Diagram (Figure 6.1).",
        "System Architecture Diagram (Figure 6.2).",
        "System Context Diagram.",
        "Overall Use Case Diagram.",
        "Entity Relationship Diagram.",
        "Component Diagram.",
        "Deployment Diagram.",
        "Activity Diagram \u2013 Student Learning Flow.",
        "Activity Diagram \u2013 Teacher Content Creation Flow.",
        "Activity Diagram \u2013 Administrator Teacher Verification Flow.",
        "Sequence Diagram \u2013 User Login.",
        "Sequence Diagram \u2013 AI Chat Tutor Request.",
        "Sequence Diagram \u2013 Quiz Attempt.",
        "Sequence Diagram \u2013 Teacher Content Creation.",
        "Sequence Diagram \u2013 Admin Teacher Verification.",
        "State Machine Diagram \u2013 Teacher Verification Status.",
        "Site Map / Navigation Flow Diagram.",
        "Selected User Interface Wireframes (Home, Student Dashboard, "
        "Teacher Content Editor, Admin Teacher Verification).",
        "Gantt Chart (Figure 8.1).",
        "Work Breakdown Structure Diagram / Table (Table 8.1).",
        "Quiz Scoring and Reward Awarding Flowchart.",
    ]
    for diagram in diagrams:
        add_number(doc, diagram)
    add_paragraph(
        doc,
        "These diagrams support the explanation of system boundaries, "
        "user roles, database structure, component interaction, deployment "
        "structure, learning flows, administrative workflows, and selected "
        "business logic in the final project documentation.",
    )

    add_heading(doc, "Appendix D. Traceability Matrix", 2)
    add_paragraph(
        doc,
        "Table D.1 provides a traceability matrix linking specific objectives "
        "to system modules and evaluation evidence. This matrix will be "
        "used during prototype development and final reporting to confirm "
        "that planned functions have been implemented and evaluated.",
    )
    add_table_caption(
        doc,
        "Table D.1. Traceability matrix linking specific objectives to "
        "modules and evaluation evidence.",
    )
    trace_rows = [
        ["SO1",
         "Requirement analysis (no system module)",
         "Literature review section, competitor matrix, expert "
         "consultation memo."],
        ["SO2",
         "All modules \u2013 design phase artefacts",
         "Architecture, ER, use case, component, deployment, activity, "
         "sequence, state diagrams; wireframes."],
        ["SO3",
         "Authentication and RBAC; Teacher Verification; Content "
         "Management; Multilingual Interface",
         "Working prototype, JUnit/MockMvc tests for auth, RBAC, "
         "verification; Vitest tests for i18n coverage."],
        ["SO4",
         "Interactive Learning Activities; Learning Content Management",
         "Implemented activities; activity records; content seed data."],
        ["SO5",
         "AI-Assisted Learning Support; Teacher Feedback and "
         "Drafting",
         "Implemented AI features; teacher review workflow; usage logs."],
        ["SO6",
         "All modules \u2013 evaluation phase",
         "SUS-equivalent score; task completion rate; perceived learning "
         "support score; expert review notes."],
    ]
    add_table(
        doc,
        header=[
            "Specific objective",
            "Mapped module(s)",
            "Evaluation evidence",
        ],
        rows=trace_rows,
        column_widths_cm=[2.0, 6.5, 7.5],
    )


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------

def build_document() -> None:
    doc = Document()
    _configure_styles(doc)
    _configure_section(doc)

    write_title_page(doc)
    write_declaration_page(doc)
    write_acknowledgement(doc)
    write_abstract(doc)
    write_table_of_contents(doc)
    write_lists_and_abbreviations(doc)
    write_introduction(doc)
    write_problem_statement(doc)
    write_objectives(doc)
    write_scope(doc)
    write_literature_review(doc)
    write_methodology(doc)
    write_roi_rov(doc)
    write_project_planning(doc)
    write_conclusion(doc)
    write_references(doc)
    write_appendices(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Document written to: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
